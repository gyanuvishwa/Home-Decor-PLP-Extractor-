#!/usr/bin/env python3
"""Build Bathroom.xlsx from the bathroom worker JSON (b<SR>.json).

Writes a BRAND NEW workbook. Never opens, reads or touches:
  Kitchen_and_Dining.xlsx / Lighting.xlsx / Wall_Decor.xlsx /
  Decorative_Home_Accessories.xlsx / Furniture.xlsx / Textile.xlsx /
  Storage.xlsx / Pet_Care.xlsx / Seasonal.xlsx / Final_Company_List (1).xlsx
  (read-only for scope + cross-category exclusion index)

Sheets written:
  Output            - same 10-column layout/format as the other category files
  Processing Ledger - one row per roster company (286)
  Review            - leaf rows with null qty, a flag, an explicit review[]
                       entry from the worker, or a cross-category dedup hit

CROSS-CATEGORY DEDUP FIX (2026-09-05): the run delivered from the other
machine only checked new leaves against Furniture.xlsx ("OPTION 1" policy in
CHECKPOINT.json). An audit found 182 leaf URLs already recorded in
Storage.xlsx (120), Wall_Decor.xlsx (61) and Textile.xlsx (1) -- the same
policy, incompletely applied. This merge extends the SAME declared policy
(never duplicate a leaf URL that another category workbook already owns) to
the full set of protected workbooks via cross_category_url_index.json (built
by build_cross_category_index.py). A hit is moved from rows[] to review[]
with an explicit issue string; nothing in Storage/Wall_Decor/Textile/etc. is
touched. This applies uniformly to the reconstructed SR 1-196 rows and every
new SR 197-286 row -- see qa_notes.md for the exact before/after counts.

Idempotent: rebuilds everything from whatever b*.json is on disk.

Usage:  python merge_bathroom.py
"""
import io, json, os, re, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

SCRATCH = os.path.dirname(os.path.abspath(__file__))
PROJ = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
BOOK = os.path.join(PROJ, "Bathroom.xlsx")
SRC = os.path.join(PROJ, "Final_Company_List (1).xlsx")
CROSS_INDEX_PATH = os.path.join(SCRATCH, "cross_category_url_index.json")

HDR = ["Maisons", "Company name", "Brand Site", "Country", "site URL",
       "Category", "Sub-Category", "Type", "qty", "link"]

HDR_FONT = Font(name="Calibri", sz=10, bold=True)
HDR_FILL = PatternFill("solid", fgColor="FFD9E1F2")
PARENT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
PARENT_FONT = Font(name="Calibri", sz=10, bold=True)
BODY_FONT = Font(name="Calibri", sz=10)
THIN = Side(style="thin", color="FFBFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BANNED = re.compile(
    r'^\s*(shop\s?all|view\s?all|browse\s?all|explore\s?all|see\s?all|all\s+products?|'
    r'shop\s+the\s+collection|discover|collections?|featured|'
    r'new\s?arrivals?|new\s?in|best\s?sellers?|sale|offers?|clearance|outlet|'
    r'gift(s)?|gift\s?guides?|gift\s?cards?|lookbooks?|blog(s)?|editorial|journal|'
    r'magazine|inspiration|ideas|stories|landing\s?pages?|'
    r'trending|recently\s?viewed|search\s+results?)\s*$', re.I)


def norm_url(u):
    return (u or "").strip().rstrip("/")


def load_cross_index():
    if not os.path.exists(CROSS_INDEX_PATH):
        return {}
    return json.load(io.open(CROSS_INDEX_PATH, encoding="utf-8"))


def load_scope():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Output"]
    scope = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            scope.append({"sr": row[0], "company": row[1], "brand_site": row[2],
                          "country": row[3], "site_url": row[4]})
    wb.close()
    return scope


def load_blocks(scope, cross_index):
    blocks, problems = [], []
    for c in scope:
        sr = c["sr"]
        p = os.path.join(SCRATCH, f"b{sr}.json")
        if not os.path.exists(p):
            blocks.append({**c, "status": "pending", "failure_reason": None,
                           "notes": None, "rows": [], "extra_review": []})
            continue
        try:
            d = json.load(io.open(p, encoding="utf-8"))
        except Exception as e:
            blocks.append({**c, "status": "error", "failure_reason":
                           f"worker JSON did not parse: {e}", "notes": None,
                           "rows": [], "extra_review": []})
            problems.append(f"SR {sr}: BAD JSON: {e}")
            continue

        rows, seen_links, seen_names = [], set(), set()
        cross_dupes = []
        for r in d.get("rows") or []:
            sub = (r.get("sub_category") or "").strip()
            cat = (r.get("category") or "").strip() or "Bathroom"
            link = (r.get("link") or "").strip()
            qty = r.get("qty")
            is_parent = bool(r.get("is_group") or r.get("parent"))
            flag = r.get("flag")

            if not sub:
                problems.append(f"SR {sr}: row with empty sub_category dropped")
                continue
            if BANNED.match(sub):
                problems.append(f"SR {sr}: dropped nav/marketing node '{sub}'")
                continue

            if not is_parent:
                hit = cross_index.get(norm_url(link)) if link else None
                if hit:
                    cross_dupes.append((cat, sub, qty, link, hit))
                    problems.append(f"SR {sr}: '{sub}' -> {link} already extracted "
                                     f"in {hit} -> moved to review")
                    continue

            if is_parent:
                if qty is not None or link:
                    problems.append(f"SR {sr}: grouping row '{sub}' had qty/link -> cleared")
                qty, link = None, ""
            else:
                key = (cat.lower(), sub.lower(), link.lower())
                if key in seen_names:
                    problems.append(f"SR {sr}: dropped exact duplicate '{cat}/{sub}' -> {link}")
                    continue
                if link and link in seen_links:
                    problems.append(f"SR {sr}: '{sub}' shares a link with another row (kept, flagged)")
                    flag = ((flag + " | ") if flag else "") + \
                        "SAME LINK AS ANOTHER ROW IN THIS COMPANY - verify not an accidental duplicate"
                if qty == 0:
                    problems.append(f"SR {sr}: dropped empty (0-product) category '{sub}'")
                    continue
                if qty is not None:
                    if not isinstance(qty, int) or isinstance(qty, bool) or qty < 0:
                        problems.append(f"SR {sr}: non-integer qty {qty!r} for '{sub}' -> nulled")
                        qty = None
                if link:
                    if not re.match(r'^https?://', link):
                        problems.append(f"SR {sr}: non-absolute link for '{sub}' -> nulled: {link}")
                        link = ""
                    elif re.search(r'(r\.jina\.ai|translate\.goog|web\.archive\.org|webcache)', link):
                        problems.append(f"SR {sr}: PROXY URL in link for '{sub}' -> nulled: {link}")
                        link = ""
                if link:
                    seen_links.add(link)
                seen_names.add(key)

            for fld, val in (("category", cat), ("sub_category", sub)):
                bad = sorted({ch for ch in val if ord(ch) > 127})
                if bad:
                    problems.append(f"SR {sr}: NON-ENGLISH chars {bad!r} in {fld} '{val}'")

            rows.append({"category": cat, "sub_category": sub, "qty": qty,
                         "link": link, "parent": is_parent,
                         "flag": flag, "evidence": r.get("evidence")})

        # Drop grouping rows left with no surviving children.
        pruned = []
        for i, row in enumerate(rows):
            if row["parent"]:
                j = i + 1
                while j < len(rows) and rows[j]["parent"]:
                    j += 1
                if j >= len(rows):
                    problems.append(f"SR {sr}: grouping row '{row['sub_category']}' "
                                    f"has no surviving children -> dropped")
                    continue
            pruned.append(row)

        extra_review = []
        for rv in d.get("review") or []:
            extra_review.append((sr, c["company"], rv.get("category"),
                                  rv.get("sub_category"), rv.get("qty"),
                                  rv.get("issue") or "flagged by worker",
                                  rv.get("link")))
        for cat, sub, qty, link, hit in cross_dupes:
            extra_review.append((sr, c["company"], cat, sub, qty,
                                  f"ALREADY EXTRACTED IN {hit} "
                                  f"(cross-category dedup fix applied at merge 2026-09-05)",
                                  link))

        blocks.append({**c,
                       "status": d.get("status") or "unknown",
                       "failure_reason": d.get("failure_reason"),
                       "notes": d.get("notes"),
                       "rows": pruned,
                       "extra_review": extra_review})
    return blocks, problems


def style(ws, r, ncols=10, parent=False):
    for c in range(1, ncols + 1):
        cell = ws.cell(r, c)
        cell.font = PARENT_FONT if parent else BODY_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = CELL_BORDER
        if parent:
            cell.fill = PARENT_FILL


def ledger_status(b):
    st = (b["status"] or "").lower()
    if st == "pending":
        return "PENDING"
    if st in ("blocked", "failed"):
        return "BLOCKED / ACCESS FAILURE"
    if st in ("error", "missing", "unknown"):
        return "PROCESSING ERROR"
    if any(not x["parent"] for x in b["rows"]):
        return "BATHROOM FOUND"
    if st == "partial":
        return "PARTIAL - UNRESOLVED (see review)"
    return "NO BATHROOM FOUND"


def main():
    lock = os.path.join(PROJ, "~$Bathroom.xlsx")
    if os.path.exists(lock):
        sys.exit(f"ABORT: Bathroom.xlsx is open in Excel (lock {lock}). Close it first.")

    cross_index = load_cross_index()
    scope = load_scope()
    assert len(scope) == 286, f"scope is {len(scope)}, expected 286"
    blocks, problems = load_blocks(scope, cross_index)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Output"

    ws.append(HDR)
    for c in range(1, 11):
        cell = ws.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 20

    r = 2
    leaf_total = parent_total = 0
    review = []
    for b in blocks:
        review.extend(b.get("extra_review") or [])
        if not b["rows"]:
            continue
        rows = b["rows"]
        prev_cat = None
        for i, row in enumerate(rows):
            if i == 0:
                ws.cell(r, 1, b["sr"]); ws.cell(r, 2, b["company"])
                ws.cell(r, 3, b["brand_site"]); ws.cell(r, 4, b["country"])
                ws.cell(r, 5, b["site_url"]); ws.cell(r, 8, "category")
            if row["category"] and row["category"] != prev_cat:
                ws.cell(r, 6, row["category"])
                prev_cat = row["category"]
            ws.cell(r, 7, row["sub_category"])
            if row["parent"]:
                parent_total += 1
            else:
                leaf_total += 1
                if row["qty"] is not None:
                    ws.cell(r, 9, row["qty"])
                if row["link"]:
                    ws.cell(r, 10, row["link"])
                if row["qty"] is None or row.get("flag"):
                    review.append((b["sr"], b["company"], row["category"],
                                   row["sub_category"], row["qty"],
                                   row.get("flag") or "qty could not be verified",
                                   row["link"]))
            style(ws, r, parent=row["parent"])
            r += 1
        ws.cell(r, 1, None); style(ws, r); r += 1

    for col, w in zip("ABCDEFGHIJ", (8, 24, 20, 10, 34, 20, 34, 10, 8, 78)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    # ---------------- Processing Ledger (all 286) ----------------
    lw = wb.create_sheet("Processing Ledger")
    LHDR = ["Maisons", "Company", "Country", "Website", "Bathroom Found",
            "Bathroom Record Count", "Status", "Notes"]
    lw.append(LHDR)
    for c in range(1, len(LHDR) + 1):
        cell = lw.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    counts = {"BATHROOM FOUND": 0, "NO BATHROOM FOUND": 0,
              "BLOCKED / ACCESS FAILURE": 0, "PROCESSING ERROR": 0, "PENDING": 0,
              "PARTIAL - UNRESOLVED (see review)": 0}
    lr = 2
    for b in blocks:
        st = ledger_status(b)
        counts[st] += 1
        leaves = [x for x in b["rows"] if not x["parent"]]
        note = b.get("failure_reason") or b.get("notes") or ""
        lw.cell(lr, 1, b["sr"]); lw.cell(lr, 2, b["company"])
        lw.cell(lr, 3, b["country"]); lw.cell(lr, 4, b["site_url"])
        lw.cell(lr, 5, "YES" if leaves else ("PENDING" if st == "PENDING" else
                       ("PARTIAL" if st.startswith("PARTIAL") else "NO")))
        lw.cell(lr, 6, len(leaves))
        lw.cell(lr, 7, st)
        lw.cell(lr, 8, str(note)[:2000])
        style(lw, lr, ncols=len(LHDR))
        lw.cell(lr, 8).alignment = Alignment(vertical="top", wrap_text=True)
        lr += 1
    for col, w in zip("ABCDEFGH", (8, 28, 12, 40, 15, 10, 26, 100)):
        lw.column_dimensions[col].width = w
    lw.freeze_panes = "A2"

    # ---------------- Review sheet ----------------
    rw = wb.create_sheet("Review")
    RHDR = ["Maisons", "Company", "Category", "Sub-Category", "qty", "Issue", "link"]
    rw.append(RHDR)
    for c in range(1, len(RHDR) + 1):
        cell = rw.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    for i, rec in enumerate(review, start=2):
        for j, v in enumerate(rec, start=1):
            rw.cell(i, j, v)
        style(rw, i, ncols=len(RHDR))
    for col, w in zip("ABCDEFG", (8, 26, 20, 32, 8, 60, 70)):
        rw.column_dimensions[col].width = w
    rw.freeze_panes = "A2"

    wb.save(BOOK)

    print(f"WROTE -> {BOOK}")
    print(f"  Output sheet: {r-1} rows | leaf rows {leaf_total} | grouping rows {parent_total}")
    print(f"  companies with rows: {sum(1 for b in blocks if any(not x['parent'] for x in b['rows']))}")
    print(f"  ledger rows: {lr-2}  (must be 286)")
    print(f"  review rows: {len(review)}")
    print()
    for k, v in counts.items():
        print(f"  {k:<26} {v}")
    print(f"  {'TOTAL':<26} {sum(counts.values())}")
    if problems:
        print(f"\nDATA-HYGIENE MESSAGES ({len(problems)}):")
        for p in problems:
            print("  " + p)


if __name__ == "__main__":
    main()
