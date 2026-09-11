#!/usr/bin/env python3
"""Post-merge validation of Bathroom.xlsx + protected-file check."""
import os, re, sys, json, hashlib, collections
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

SCRATCH = os.path.dirname(os.path.abspath(__file__))
PROJ = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
BOOK = os.path.join(PROJ, "Bathroom.xlsx")
PROTECTED = ["Kitchen_and_Dining.xlsx", "Lighting.xlsx", "Wall_Decor.xlsx",
             "Decorative_Home_Accessories.xlsx", "Furniture.xlsx", "Textile.xlsx",
             "Storage.xlsx", "Pet_Care.xlsx", "Seasonal.xlsx",
             "Final_Company_List (1).xlsx"]

BANNED_URL = re.compile(r'(r\.jina\.ai|translate\.goog|web\.archive\.org|/search|[?&]q=)', re.I)


def main():
    print("=== PROTECTED FILES ===")
    base = {}
    bl = os.path.join(SCRATCH, "protected_baseline.sha256")
    if os.path.exists(bl):
        for line in open(bl, encoding="utf-8"):
            h, _, name = line.strip().partition(" ")
            base[name.lstrip("*").strip()] = h
    for f in PROTECTED:
        p = os.path.join(PROJ, f)
        if not os.path.exists(p):
            print(f"  {f:<38} (not found)")
            continue
        h = hashlib.sha256(open(p, "rb").read()).hexdigest()
        exp = base.get(f)
        state = "UNCHANGED" if exp == h else ("MODIFIED !!!" if exp else "(no baseline)")
        print(f"  {f:<38} {state}")

    print("\n=== Bathroom.xlsx ===")
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    print("  sheets:", wb.sheetnames)
    ws = wb["Output"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    companies = [r for r in rows if r[0] is not None]
    leaves = [r for r in rows if (r[8] is not None or r[9])]
    groups = [r for r in rows if r[6] and r[8] is None and not r[9]]
    blanks = [r for r in rows if not any(v is not None for v in r)]
    print(f"  total rows          {len(rows)}")
    print(f"  companies           {len(companies)}")
    print(f"  leaf rows           {len(leaves)}")
    print(f"  grouping rows       {len(groups)}")
    print(f"  separator rows      {len(blanks)}")

    srs = [r[0] for r in companies]
    assert len(srs) == len(set(srs)), "DUPLICATE SR in Output"
    print(f"  SR unique           yes ({min(srs)}..{max(srs)})" if srs else "  (empty)")

    # duplicate link within the whole Bathroom Output (own-file dedup)
    all_links = [r[9] for r in rows if r[9]]
    own_dupes = {k: v for k, v in collections.Counter(all_links).items() if v > 1}
    print(f"  duplicate links WITHIN Bathroom Output: {len(own_dupes)}")
    for k, v in list(own_dupes.items())[:20]:
        print("      ", v, k)

    # STRICT: same URL used by >1 company within Bathroom itself
    cur = None
    link_owners = collections.defaultdict(set)
    for r in rows:
        if r[0] is not None:
            cur = (r[0], r[1])
        if r[9]:
            link_owners[r[9]].add(cur)
    cross_company_dupes = {link: owners for link, owners in link_owners.items() if len(owners) > 1}
    print(f"  STRICT: same URL used by >1 company (within Bathroom): {len(cross_company_dupes)}")
    for link, owners in list(cross_company_dupes.items())[:20]:
        print("      ", link, "->", owners)

    # cross-category contamination: any Bathroom leaf URL that ALSO exists in
    # another protected category workbook's Output. Must be zero after the
    # 2026-09-05 dedup fix (merge_bathroom.py already routes hits to review[]
    # at merge time using cross_category_url_index.json -- this is the
    # after-the-fact proof it actually worked).
    idx_path = os.path.join(SCRATCH, "cross_category_url_index.json")
    cross_hits = []
    if os.path.exists(idx_path):
        cross_index = json.load(open(idx_path, encoding="utf-8"))
        for r in rows:
            link = r[9]
            if link:
                owner = cross_index.get(link.strip().rstrip("/"))
                if owner:
                    cross_hits.append((r[1], r[6], link, owner))
    print(f"  CROSS-CATEGORY contamination (Bathroom leaf already in another workbook): {len(cross_hits)}")
    for x in cross_hits[:20]:
        print("      ", x)

    # non-ASCII
    na = [(r[5], r[6]) for r in rows if any(ord(ch) > 127 for ch in str(r[5] or "") + str(r[6] or ""))]
    print(f"  non-ASCII cells     {len(na)}")
    for x in na[:20]:
        print("      ", x)

    # bad URLs
    bad = [(r[6], r[9]) for r in rows if r[9] and (BANNED_URL.search(str(r[9])) or not str(r[9]).startswith("http"))]
    print(f"  suspect URLs        {len(bad)}")
    for x in bad[:20]:
        print("      ", x)

    q = [r[8] for r in leaves if isinstance(r[8], int)]
    print(f"  leaf rows with qty  {len(q)} / {len(leaves)}   (null qty {len(leaves)-len(q)})")

    lw = wb["Processing Ledger"]
    lrows = list(lw.iter_rows(min_row=2, values_only=True))
    lrows = [r for r in lrows if r[0] is not None]
    st = collections.Counter(r[6] for r in lrows)
    print("\n=== LEDGER ===")
    print(f"  ledger rows         {len(lrows)}  (must be 286)")
    lsrs = [r[0] for r in lrows]
    print(f"  unique SR           {len(set(lsrs))}")
    for k, v in st.items():
        print(f"  {k:<26} {v}")
    print(f"  {'SUM':<26} {sum(st.values())}")
    assert len(lrows) == 286, f"LEDGER IS NOT 286 ROWS (got {len(lrows)})"

    rw = wb["Review"]
    rrows = [r for r in rw.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    print(f"\n  review rows         {len(rrows)}")

    ok = sum(st.values()) == 286 and not cross_company_dupes and not own_dupes and not cross_hits and not na
    print("\nOK" if ok else "\nRECONCILIATION FAILED")
    assert not cross_company_dupes, f"STRICT DUPLICATE URL CHECK FAILED -- {len(cross_company_dupes)} link(s)"
    assert not own_dupes, f"DUPLICATE LINK WITHIN OUTPUT -- {len(own_dupes)} link(s)"
    assert not cross_hits, f"CROSS-CATEGORY CONTAMINATION -- {len(cross_hits)} link(s) still leaked into rows"
    assert not na, f"NON-ASCII CELLS REMAIN -- {len(na)}"


if __name__ == "__main__":
    main()
