import sys, collections, re
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

BOOK = r"C:\Users\GyanendraVishwakarma\Web Research Agent\Final_Company_List (1).xlsx"
wb = openpyxl.load_workbook(BOOK, data_only=True)
ws = wb["Output"]

links, nonascii, badlink, qtybad = collections.Counter(), [], [], []
leaf = parent = blank = 0
srs = set()
for r in range(2, ws.max_row + 1):
    sn, cat, sub = ws.cell(r, 1).value, ws.cell(r, 6).value, ws.cell(r, 7).value
    qty, link = ws.cell(r, 9).value, ws.cell(r, 10).value
    if isinstance(sn, int):
        srs.add(sn)
    if sub is None and qty is None and link is None and cat is None:
        blank += 1
        continue
    fill = ws.cell(r, 7).fill
    is_par = fill and fill.fgColor and fill.fgColor.rgb == "FFFFF2CC"
    if is_par:
        parent += 1
        if qty is not None or link:
            qtybad.append((r, sub, qty, link))
    else:
        leaf += 1
        if link:
            links[link] += 1
            if not re.match(r"^https?://", str(link)):
                badlink.append((r, link))
            if re.search(r"(translate\.goog|r\.jina\.ai|web\.archive\.org)", str(link)):
                badlink.append((r, link))
    for v in (cat, sub):
        if v and any(ord(ch) > 127 for ch in str(v)):
            nonascii.append((r, v))

dups = {k: v for k, v in links.items() if v > 1}
print("Output rows          :", ws.max_row - 1)
print("  leaf / parent      : %d / %d   (separators %d)" % (leaf, parent, blank))
print("  distinct S.N.      :", len(srs), "->", "contiguous 1-30" if srs == set(range(1, 31)) else sorted(srs))
print("  unique leaf links  :", len(links))
print("  DUPLICATE links    :", len(dups))
for k, v in list(dups.items())[:10]:
    print("     x%d %s" % (v, k))
print("  non-ASCII cells    :", len(nonascii), nonascii[:5])
print("  proxy/relative link:", len(badlink), badlink[:5])
print("  parents w/ qty|link:", len(qtybad), qtybad[:5])

mws = wb["Master Company List"]
st = collections.Counter(mws.cell(r, 5).value for r in range(2, 32))
print("\nMaster Status (S.N. 1-30):", dict(st))
blanks = [mws.cell(r, 1).value for r in range(32, mws.max_row + 1) if mws.cell(r, 5).value]
print("Rows past S.N. 30 with a Status:", blanks or "none (correct)")
