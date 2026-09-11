#!/usr/bin/env python3
"""Build Textile.xlsx and Storage.xlsx from the combined worker JSON (ts<SR>.json).

One shared research pass per company (ts<SR>.json) is split into two
independent workbooks, each with the same 10-column layout/format used by
Furniture.xlsx and the four decor category workbooks.

Never opens, reads or touches any other validated workbook.

Sheets written per workbook:
  Output            - same 10-column layout/format as the other category files
  Processing Ledger - one row per company in scope (must match roster size)
  Review            - leaf rows with null qty or a MANUAL REVIEW / LOW-confidence flag

Idempotent: rebuilds everything from whatever ts*.json is on disk.

Usage:  python merge_textile_storage.py
"""
import io, json, os, re, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

SCRATCH = os.path.dirname(os.path.abspath(__file__))
PROJ = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
ROSTER = os.path.join(SCRATCH, "companies.json")

HDR = ["Maisons", "Company name", "Brand Site", "Country", "site URL",
       "Category", "Sub-Category", "Type", "qty", "link"]

HDR_FONT = Font(name="Calibri", sz=10, bold=True)
HDR_FILL = PatternFill("solid", fgColor="FFD9E1F2")
PARENT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
PARENT_FONT = Font(name="Calibri", sz=10, bold=True)
BODY_FONT = Font(name="Calibri", sz=10)
THIN = Side(style="thin", color="FFBFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Navigation / marketing nodes that are never a product category.
BANNED = re.compile(
    r'^\s*(shop\s?all|view\s?all|browse\s?all|explore\s?all|see\s?all|all\s+products?|'
    r'shop\s+the\s+collection|discover|collections?|featured|'
    r'new\s?arrivals?|new\s?in|best\s?sellers?|sale|offers?|clearance|outlet|'
    r'gift\s?guide|gift\s?cards?|lookbook|blog|journal|magazine|inspiration|ideas|'
    r'stories|designers?|brands?|trending|recently\s?viewed|search\s+results?)\s*$', re.I)

CATEGORIES = {
    "textile": {"label": "Textile", "book": os.path.join(PROJ, "Textile.xlsx"),
                "field": "textile_rows"},
    "storage": {"label": "Storage", "book": os.path.join(PROJ, "Storage.xlsx"),
                "field": "storage_rows"},
}


def load_roster():
    d = json.load(io.open(ROSTER, encoding="utf-8"))
    scope = []
    for c in d:
        scope.append({"sr": c["sr"], "company": c["name"], "brand_site": c["brand_site"],
                      "country": c["country"], "site_url": c["url"]})
    return scope


def load_worker_docs(scope):
    docs, problems = {}, []
    for c in scope:
        sr = c["sr"]
        p = os.path.join(SCRATCH, f"ts{sr}.json")
        if not os.path.exists(p):
            docs[sr] = {**c, "status": "MISSING", "failure_reason":
                        f"no ts{sr}.json produced", "notes": None,
                        "textile_rows": [], "storage_rows": []}
            problems.append(f"SR {sr}: NO FILE (ts{sr}.json missing)")
            continue
        try:
            d = json.load(io.open(p, encoding="utf-8"))
        except Exception as e:
            docs[sr] = {**c, "status": "error", "failure_reason":
                        f"worker JSON did not parse: {e}", "notes": None,
                        "textile_rows": [], "storage_rows": []}
            problems.append(f"SR {sr}: BAD JSON: {e}")
            continue
        docs[sr] = {**c, "status": d.get("status") or "unknown",
                    "failure_reason": d.get("failure_reason"), "notes": d.get("notes"),
                    "textile_rows": d.get("textile_rows") or [],
                    "storage_rows": d.get("storage_rows") or []}
    return docs, problems


def clean_rows(sr, raw_rows, label, problems):
    rows, seen_links, seen_names = [], set(), set()
    for r in raw_rows:
        sub = (r.get("sub_category") or "").strip()
        cat = (r.get("category") or "").strip() or label
        link = (r.get("link") or "").strip()
        qty = r.get("qty")
        is_parent = bool(r.get("is_group") or r.get("parent"))

        if not sub:
            problems.append(f"SR {sr} [{label}]: row with empty sub_category dropped")
            continue
        if BANNED.match(sub):
            problems.append(f"SR {sr} [{label}]: dropped nav/marketing node '{sub}'")
            continue

        if is_parent:
            if qty is not None or link:
                problems.append(f"SR {sr} [{label}]: grouping row '{sub}' had qty/link -> cleared")
            qty, link = None, ""
        else:
            key = (cat.lower(), sub.lower(), link.lower())
            if key in seen_names:
                problems.append(f"SR {sr} [{label}]: dropped exact duplicate '{cat}/{sub}' -> {link}")
                continue
            if link and link in seen_links:
                problems.append(f"SR {sr} [{label}]: '{sub}' shares a link with another row (kept, flagged)")
                r = dict(r)
                r["flag"] = ((r.get("flag") + " | ") if r.get("flag") else "") + \
                    "SAME LINK AS ANOTHER ROW IN THIS COMPANY - verify not an accidental duplicate"
            if qty == 0:
                pass  # explicit zero-product category is valid per rule files §39/§34 (storage) - keep
            if qty is not None:
                if not isinstance(qty, int) or isinstance(qty, bool) or qty < 0:
                    problems.append(f"SR {sr} [{label}]: non-integer qty {qty!r} for '{sub}' -> nulled")
                    qty = None
                elif not (r.get("evidence") or "").strip():
                    problems.append(f"SR {sr} [{label}]: qty {qty} for '{sub}' has NO evidence -> nulled")
                    qty = None
            if link:
                if not re.match(r'^https?://', link):
                    problems.append(f"SR {sr} [{label}]: non-absolute link for '{sub}' -> nulled: {link}")
                    link = ""
                elif re.search(r'(r\.jina\.ai|translate\.goog|web\.archive\.org|webcache)', link):
                    problems.append(f"SR {sr} [{label}]: PROXY URL in link for '{sub}' -> nulled: {link}")
                    link = ""
                elif re.search(r'[?&](q|query|search|keyword)=', link) or '/search' in link.lower():
                    problems.append(f"SR {sr} [{label}]: SEARCH URL in link for '{sub}' (kept, flagged): {link}")
                    r = dict(r)
                    r["flag"] = ((r.get("flag") + " | ") if r.get("flag") else "") + \
                        "LINK LOOKS LIKE A SEARCH PAGE - verify"
            if link:
                seen_links.add(link)
            seen_names.add(key)

        for fld, val in (("category", cat), ("sub_category", sub)):
            bad = sorted({ch for ch in val if ord(ch) > 127})
            if bad:
                problems.append(f"SR {sr} [{label}]: NON-ENGLISH chars {bad!r} in {fld} '{val}'")

        flag = r.get("flag")
        conf = (r.get("confidence") or "").upper()
        if conf == "LOW" and not is_parent:
            flag = ((flag + " | ") if flag else "") + "LOW CONFIDENCE - review queue"

        rows.append({"category": cat, "sub_category": sub, "qty": qty,
                     "link": link, "parent": is_parent,
                     "flag": flag, "evidence": r.get("evidence")})

    pruned = []
    for i, row in enumerate(rows):
        if row["parent"]:
            j = i + 1
            while j < len(rows) and rows[j]["parent"]:
                j += 1
            if j >= len(rows):
                problems.append(f"SR {sr} [{label}]: grouping row '{row['sub_category']}' "
                                f"has no surviving children -> dropped")
                continue
        pruned.append(row)
    return pruned


def style(ws, r, ncols=10, parent=False):
    for c in range(1, ncols + 1):
        cell = ws.cell(r, c)
        cell.font = PARENT_FONT if parent else BODY_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = CELL_BORDER
        if parent:
            cell.fill = PARENT_FILL


def ledger_status(b, label):
    st = (b["status"] or "").lower()
    if st in ("blocked", "failed"):
        return "BLOCKED / ACCESS FAILURE"
    if st in ("error", "missing", "unknown"):
        return "PROCESSING ERROR"
    return f"{label.upper()} FOUND" if b["rows"] else f"NO {label.upper()} FOUND"


def build_workbook(key, cfg, blocks, roster_len):
    label = cfg["label"]
    book = cfg["book"]
    lock = os.path.join(PROJ, "~$" + os.path.basename(book))
    if os.path.exists(lock):
        sys.exit(f"ABORT: {os.path.basename(book)} is open in Excel (lock {lock}). Close it first.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Output"
    ws.append(HDR)
    for c in range(1, 11):
        cell = ws.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 20

    r = 2
    leaf_total = parent_total = 0
    review = []
    for b in blocks:
        if not b["rows"]:
            continue
        rows = b["rows"]
        prev_cat = None
        for i, row in enumerate(rows):
            if i == 0:
                ws.cell(r, 1, b["sr"]); ws.cell(r, 2, b["company"])
                ws.cell(r, 3, b["brand_site"]); ws.cell(r, 4, b["country"])
                ws.cell(r, 5, b["site_url"]); ws.cell(r, 8, "category")
            if row["category"] and row["category"] != prev_cat:
                ws.cell(r, 6, row["category"])
                prev_cat = row["category"]
            ws.cell(r, 7, row["sub_category"])
            if row["parent"]:
                parent_total += 1
            else:
                leaf_total += 1
                if row["qty"] is not None:
                    ws.cell(r, 9, row["qty"])
                if row["link"]:
                    ws.cell(r, 10, row["link"])
                if row["qty"] is None or row.get("flag"):
                    review.append((b["sr"], b["company"], row["category"],
                                   row["sub_category"], row["qty"],
                                   row.get("flag") or "qty could not be verified",
                                   row["link"]))
            style(ws, r, parent=row["parent"])
            r += 1
        ws.cell(r, 1, None); style(ws, r); r += 1

    for col, w in zip("ABCDEFGHIJ", (8, 24, 20, 10, 34, 20, 34, 10, 8, 78)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    lw = wb.create_sheet("Processing Ledger")
    LHDR = ["Maisons", "Company", "Country", "Website", f"{label} Found",
            f"{label} Record Count", "Status", "Notes"]
    lw.append(LHDR)
    for c in range(1, len(LHDR) + 1):
        cell = lw.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    counts = {f"{label.upper()} FOUND": 0, f"NO {label.upper()} FOUND": 0,
              "BLOCKED / ACCESS FAILURE": 0, "PROCESSING ERROR": 0}
    lr = 2
    for b in blocks:
        st = ledger_status(b, label)
        counts[st] += 1
        leaves = [x for x in b["rows"] if not x["parent"]]
        note = b.get("failure_reason") or b.get("notes") or ""
        lw.cell(lr, 1, b["sr"]); lw.cell(lr, 2, b["company"])
        lw.cell(lr, 3, b["country"]); lw.cell(lr, 4, b["site_url"])
        lw.cell(lr, 5, "YES" if b["rows"] else "NO")
        lw.cell(lr, 6, len(leaves))
        lw.cell(lr, 7, st)
        lw.cell(lr, 8, str(note)[:2000])
        style(lw, lr, ncols=len(LHDR))
        lw.cell(lr, 8).alignment = Alignment(vertical="top", wrap_text=True)
        lr += 1
    for col, w in zip("ABCDEFGH", (8, 28, 12, 40, 15, 10, 26, 100)):
        lw.column_dimensions[col].width = w
    lw.freeze_panes = "A2"

    rw = wb.create_sheet("Review")
    RHDR = ["Maisons", "Company", "Category", "Sub-Category", "qty", "Issue", "link"]
    rw.append(RHDR)
    for c in range(1, len(RHDR) + 1):
        cell = rw.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    for i, rec in enumerate(review, start=2):
        for j, v in enumerate(rec, start=1):
            rw.cell(i, j, v)
        style(rw, i, ncols=len(RHDR))
    for col, w in zip("ABCDEFG", (8, 26, 20, 32, 8, 60, 70)):
        rw.column_dimensions[col].width = w
    rw.freeze_panes = "A2"

    wb.save(book)

    print(f"WROTE -> {book}")
    print(f"  Output sheet: {r-1} rows | leaf rows {leaf_total} | grouping rows {parent_total}")
    print(f"  companies with rows: {sum(1 for b in blocks if b['rows'])}")
    print(f"  ledger rows: {lr-2}  (roster size {roster_len})")
    print(f"  review rows: {len(review)}")
    for k, v in counts.items():
        print(f"  {k:<26} {v}")
    print(f"  {'TOTAL':<26} {sum(counts.values())}")


def main():
    scope = load_roster()
    docs, problems = load_worker_docs(scope)

    for key, cfg in CATEGORIES.items():
        blocks = []
        for c in scope:
            sr = c["sr"]
            d = docs[sr]
            rows = clean_rows(sr, d[cfg["field"]], cfg["label"], problems)
            blocks.append({**c, "status": d["status"], "failure_reason": d["failure_reason"],
                          "notes": d["notes"], "rows": rows})
        print("=" * 70)
        build_workbook(key, cfg, blocks, len(scope))
        print()

    if problems:
        print(f"DATA-HYGIENE MESSAGES ({len(problems)}):")
        for p in problems:
            print("  " + p)


if __name__ == "__main__":
    main()
