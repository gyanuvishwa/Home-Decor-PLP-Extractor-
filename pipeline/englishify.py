#!/usr/bin/env python3
"""Normalise Sheet3 Category / Sub-Category text in Batch1-8 to plain English.

Three repairs, applied only to columns F (Category) and G (Sub-Category):

1. ORPHAN PARENS - the native-script half of "<native> (English)" labels is gone
   from the workbook (stripped outside this pipeline), leaving remnants like
   " (Pendant Lights)" or "(Tableware)". Unwrap to "Pendant Lights".
2. RESIDUAL CJK  - any surviving Japanese/Chinese text, incl. "English (CJK)"
   pairs, is removed; the English half is kept.
3. ACCENTS       - Latin-1/Latin-Extended letters folded to ASCII
   (Decor, Consomme, Glogg). Company names in column B are NOT touched:
   they are proper nouns, not language content.

Whitespace is trimmed and internal runs collapsed.

Writes a full change log to englishify_log.tsv. Idempotent - safe to re-run.
"""
import os, re, sys, unicodedata
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

PROJ = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
TARGET = os.path.join(PROJ, "Home_Decor_Extraction_Batch1-8.xlsx")
HERE = os.path.dirname(os.path.abspath(__file__))

CJK = re.compile(r'[\u3000-\u303f\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff'
                 r'\uf900-\ufaff\uff00-\uffef]')
# "(English)" occupying the whole cell, i.e. the native half was removed.
ORPHAN = re.compile(r'^\s*\(([^()]*)\)\s*$')


def fold(s):
    """Fold accented Latin to ASCII, leaving other characters alone."""
    out = []
    for ch in unicodedata.normalize("NFD", s):
        if unicodedata.combining(ch):
            continue
        out.append(ch)
    s = "".join(out)
    return (s.replace("\u00df", "ss").replace("\u00e6", "ae").replace("\u00c6", "AE")
             .replace("\u00f8", "o").replace("\u00d8", "O")
             .replace("\u0153", "oe").replace("\u0152", "OE"))


def clean(v):
    if not isinstance(v, str):
        return v
    s = v

    # 2. drop CJK. "English (CJK)" -> "English"; "(CJK) English" -> "English".
    if CJK.search(s):
        s = re.sub(r'\s*\([^()]*\)\s*',
                   lambda m: ' ' if CJK.search(m.group(0)) else m.group(0), s)
        s = CJK.sub('', s)

    # 1. unwrap an orphan "(English)" that is the entire cell.
    for _ in range(3):
        m = ORPHAN.match(s)
        if not m or not m.group(1).strip():
            break
        s = m.group(1)

    # 3. fold accents.
    s = fold(s)

    # 4. ASCII punctuation (em/en dash, smart quotes) - keeps the sheet plain-text.
    for a, b in (("—", "-"), ("–", "-"), ("’", "'"),
                 ("‘", "'"), ("“", '"'), ("”", '"'),
                 ("·", "-"), ("…", "...")):
        s = s.replace(a, b)

    # tidy
    s = re.sub(r'\(\s*\)', '', s)
    s = re.sub(r'\s{2,}', ' ', s).strip()
    s = re.sub(r'^[|/,;&·-]+\s*', '', s).strip()
    s = re.sub(r'\s*[|/,;&·-]+$', '', s).strip()
    return s


def main():
    wb = openpyxl.load_workbook(TARGET)
    ws = wb["Sheet3"]
    log = []
    who = None
    for rr in range(2, ws.max_row + 1):
        if ws.cell(rr, 2).value:
            who = ws.cell(rr, 2).value
        for cc in (6, 7):
            cell = ws.cell(rr, cc)
            old = cell.value
            if not isinstance(old, str):
                continue
            new = clean(old)
            if new != old:
                if not new:
                    # Whole cell was native script. A blank Category cell is the
                    # sheet's own "same as above" convention, so blank it properly
                    # rather than leaving a whitespace string behind.
                    cell.value = None
                    log.append((rr, cc, who, old, "<blanked - inherits category above>"))
                    continue
                cell.value = new
                log.append((rr, cc, who, old, new))
    wb.save(TARGET)

    p = os.path.join(HERE, "englishify_log.tsv")
    with open(p, "w", encoding="utf-8") as f:
        f.write("row\tcol\tcompany\tbefore\tafter\n")
        for r in log:
            f.write("\t".join(str(x) for x in r) + "\n")

    # verify
    wb2 = openpyxl.load_workbook(TARGET)["Sheet3"]
    left = []
    for rr in range(2, wb2.max_row + 1):
        for cc in (6, 7):
            v = wb2.cell(rr, cc).value
            if isinstance(v, str) and any(ord(ch) > 127 for ch in v):
                left.append((rr, cc, v))
    print(f"changed {len(log)} cells -> {TARGET}")
    print(f"log -> {p}")
    print(f"non-ASCII cells remaining in Category/Sub-Category: {len(left)}")
    for x in left[:10]:
        print("   ", x)
    print("\nsample changes:")
    for r in log[:12]:
        print(f"   row {r[0]:>5} {r[2][:20]:<20} {r[3]!r}  ->  {r[4]!r}")


if __name__ == "__main__":
    main()
