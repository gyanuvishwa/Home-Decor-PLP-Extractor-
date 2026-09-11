"""Apply duplicate-link repairs (fix<SR>.json) to the master Output sheet and the
four category workbooks.

Each fix file lists `replacements` against a company's rows, keyed by the row's
CURRENT Sub-Category text (col G) and, when that name is ambiguous, its link:

  {"action": "repoint", "old_sub_category": "Vases", "sub_category": "Vases",
   "qty": 32, "link": "...", "evidence": "...", "flag": null}
  {"action": "merge",   "old_sub_category": ["A", "B"], "sub_category": "A & B",
   "qty": 65, "link": "...", "evidence": "..."}
  {"action": "delete",  "old_sub_category": "Bogus Row",
   "old_link": "https://...", "reason": "..."}

`old_link` narrows a name that occurs more than once in the block -- Home Box UAE
has two rows called "Serveware", one correct and one bogus, and deleting by name
alone removed both.

A merge keeps the FIRST name listed, so the merge yields exactly one row across
the whole file set; a category workbook that does not hold that first row deletes
all of its listed rows instead of keeping a second copy.

The sheet is read into memory, edited there, and written back wholesale. Editing
in place with openpyxl's delete_rows() left stray link values on unrelated
grouping rows further down the sheet, so we never shift rows inside the sheet.

Run:  python pipeline/apply_fixes.py <scratchpad-dir> [SR ...]
"""
import glob
import json
import os
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
FILES = [
    (os.path.join(ROOT, "Final_Company_List (1).xlsx"), "Output"),
    (os.path.join(ROOT, "Kitchen_and_Dining.xlsx"), "Output"),
    (os.path.join(ROOT, "Lighting.xlsx"), "Output"),
    (os.path.join(ROOT, "Wall_Decor.xlsx"), "Output"),
    (os.path.join(ROOT, "Decorative_Home_Accessories.xlsx"), "Output"),
]

NCOL = 10
SR, NAME, CAT, SUB, QTY, LINK = 0, 1, 5, 6, 8, 9   # 0-based into a row list


class Row:
    __slots__ = ("v", "group")

    def __init__(self, v, group):
        self.v = v            # list of 10 values
        self.group = group    # True if styled as a Category grouping row

    def blank(self):
        return all(x is None for x in self.v)

    def is_group(self):
        """Category grouping row: Category text, but no qty and no link."""
        return self.v[CAT] is not None and self.v[QTY] is None and self.v[LINK] is None


def read_sheet(ws):
    rows = []
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, NCOL + 1)]
        rows.append(Row(vals, bool(ws.cell(r, CAT + 1).font.bold)))
    return rows


def write_sheet(ws, rows, group_fill, header_style):
    for i, row in enumerate(rows, 1):
        for c in range(1, NCOL + 1):
            cell = ws.cell(i, c)
            cell.value = row.v[c - 1]
            if i == 1:
                continue
            if row.group:
                cell.fill = group_fill
                cell.font = Font(bold=True)
            else:
                cell.fill = PatternFill()
                cell.font = Font()
    for r in range(len(rows) + 1, ws.max_row + 1):
        for c in range(1, NCOL + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell.fill = PatternFill()
            cell.font = Font()


def block_range(rows, sr):
    start = None
    for i, row in enumerate(rows):
        if i == 0:
            continue
        if row.v[SR] is not None:
            if start is not None:
                return start, i - 1
            if row.v[SR] == sr:
                start = i
    return (start, len(rows) - 1) if start is not None else (None, None)


def apply_fix(rows, fix, path, log):
    """Apply one fix to an in-memory row list; returns the new row list.

    NOT idempotent -- a merge whose first name is already consumed falls through
    to the delete-all branch. Each workbook must be loaded fresh and each fix
    applied exactly once.
    """
    start, end = block_range(rows, fix["sr"])
    if start is None:
        log.append(f"    {os.path.basename(path)}: company not present")
        return rows
    while end > start and rows[end].blank():
        end -= 1

    def find(name, link=None):
        """Indices in the block whose Sub-Category is `name` (and link matches)."""
        out = []
        for i in range(start, end + 1):
            s = rows[i].v[SUB]
            if s is not None and str(s).strip() == name:
                if link is None or rows[i].v[LINK] == link:
                    out.append(i)
        return out

    doomed, touched, missing = set(), 0, []

    for rep in fix["replacements"]:
        act = rep["action"]
        olds = rep["old_sub_category"]
        olds = olds if isinstance(olds, list) else [olds]
        link_filter = rep.get("old_link")
        hits = {o: find(o, link_filter) for o in olds}
        hits = {o: v for o, v in hits.items() if v}
        if not hits:
            missing.append("/".join(olds))
            continue
        if act == "delete":
            for v in hits.values():
                doomed.update(v)
            touched += 1
        elif act == "repoint":
            i = hits[olds[0]][0] if olds[0] in hits else next(iter(hits.values()))[0]
            rows[i].v[SUB] = rep["sub_category"]
            rows[i].v[QTY] = rep["qty"]
            rows[i].v[LINK] = rep["link"]
            touched += 1
        elif act == "merge":
            if olds[0] in hits:
                keep = hits[olds[0]][0]
                rows[keep].v[SUB] = rep["sub_category"]
                rows[keep].v[QTY] = rep["qty"]
                rows[keep].v[LINK] = rep["link"]
                for v in hits.values():
                    doomed.update(x for x in v if x != keep)
            else:
                for v in hits.values():
                    doomed.update(v)
            touched += 1
        else:
            raise ValueError("unknown action " + act)

    # A Category grouping row whose section lost every leaf goes too.
    for i in range(start + 1, end + 1):
        if i in doomed or not rows[i].is_group():
            continue
        j, section = i + 1, []
        while j <= end and not rows[j].is_group():
            section.append(j)
            j += 1
        if section and all(s in doomed for s in section):
            doomed.add(i)

    survivors = [i for i in range(start, end + 1) if i not in doomed]
    body = any(rows[i].v[LINK] is not None for i in survivors)
    if not body:
        doomed.update(range(start, end + 1))
        if end + 1 < len(rows) and rows[end + 1].blank():
            doomed.add(end + 1)
        log.append(f"    {os.path.basename(path)}: block emptied, removed "
                   f"{len(doomed)} row(s)")
    elif start in doomed:
        # The banner row carries the company identity in cols A-E; if it was
        # deleted, promote the first survivor so the block keeps its header.
        first = min(survivors)
        for c in range(0, 5):
            rows[first].v[c] = rows[start].v[c]
        if rows[first].v[CAT] is None:
            rows[first].v[CAT] = rows[start].v[CAT]
        rows[first].group = False

    out = [r for i, r in enumerate(rows) if i not in doomed]

    if body:
        note = (f"    {os.path.basename(path)}: {touched} action(s), "
                f"{len(doomed)} row(s) deleted")
        if missing:
            note += f"  [not in this file: {', '.join(missing)}]"
        log.append(note)
    return out


def main():
    scratch = sys.argv[1]
    only = [int(x) for x in sys.argv[2:]] or None
    files = {json.load(open(f, encoding="utf-8"))["sr"]: f
             for f in glob.glob(os.path.join(scratch, "fix*.json"))}
    order = only if only else sorted(files)
    fixes = []
    for sr in order:
        fix = json.load(open(files[sr], encoding="utf-8"))
        if fix.get("status") == "failed":
            print(f"SR{sr} {fix['company']}: SKIPPED (status failed)")
            continue
        fixes.append(fix)

    # One load/save per workbook, every fix applied to it exactly once.
    group_fill = PatternFill("solid", fgColor="FFFFF2CC")
    for path, sheet in FILES:
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet]
        rows = read_sheet(ws)
        print(f"=== {os.path.basename(path)}")
        for fix in fixes:
            log = []
            rows = apply_fix(rows, fix, path, log)
            for line in log:
                line = line.strip().replace(os.path.basename(path) + ":", "").strip()
                print(f"   SR{fix['sr']:<4} {fix['company']:<18} {line}")
        write_sheet(ws, rows, group_fill, None)
        wb.save(path)


if __name__ == "__main__":
    main()
