#!/usr/bin/env python3
"""Build New_Arrivals.xlsx from the new-arrivals worker JSON (na<SR>.json).

Writes a BRAND NEW workbook. Never opens, reads or touches:
  Kitchen_and_Dining.xlsx / Lighting.xlsx / Wall_Decor.xlsx /
  Decorative_Home_Accessories.xlsx / Furniture.xlsx / Textile.xlsx /
  Storage.xlsx / Pet_Care.xlsx / Seasonal.xlsx / Bathroom.xlsx /
  Final_Company_List (1).xlsx / Bestsellers.xlsx / Clearance.xlsx
  (read-only for scope; the last two are the sibling merchandising-PLP
  deliverables and must never be written by this merge)

Sheets written:
  Output            - same 10-column layout/format as the other category files
  Processing Ledger - one row per company in scope (must be exactly 285,
                       286 roster entries minus the known SR 286/165 duplicate)
  Review            - leaf rows with null qty, a MANUAL REVIEW flag, or an
                       unresolved "New Collection"-style editorial ambiguity

This is PLP-LEVEL extraction, not a product taxonomy. Every row's category is
"New Arrivals" (or a genuine site-defined child of it, e.g. "Furniture New
Arrivals"); sub_category is null unless the site genuinely exposes a
sub-listing. Trusts the worker's PLP-vs-marketing judgement
(rules/new-arrivals.md) with ONE mandatory exception this merge enforces
itself: rules/new-arrivals.md SS8/SS10 - a bare "New Collection"/"New Season"/
"New Range" name, with none of the unambiguous New Arrivals signal words
alongside it, is force-flagged for manual review rather than trusted, the
same "force-flag an open policy question rather than trust the worker"
pattern already used by merge_pet_care.py and merge_clearance.py.

Idempotent: rebuilds everything from whatever na*.json is on disk.

Usage:  python merge_new_arrivals.py
"""
import io, json, os, re, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

SCRATCH = os.path.dirname(os.path.abspath(__file__))
PROJ = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
BOOK = os.path.join(PROJ, "New_Arrivals.xlsx")
SRC = os.path.join(PROJ, "Final_Company_List (1).xlsx")

HDR = ["Maisons", "Company name", "Brand Site", "Country", "site URL",
       "Category", "Sub-Category", "Type", "qty", "link"]

HDR_FONT = Font(name="Calibri", sz=10, bold=True)
HDR_FILL = PatternFill("solid", fgColor="FFD9E1F2")
PARENT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
PARENT_FONT = Font(name="Calibri", sz=10, bold=True)
BODY_FONT = Font(name="Calibri", sz=10)
THIN = Side(style="thin", color="FFBFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Navigation/marketing nodes that are never a genuine New Arrivals PLP, even if
# discovered while looking for one (rules/new-arrivals.md SS9). Deliberately
# does NOT ban "new arrivals/new in/just in/what's new/latest/recently added"
# etc. - those ARE the target concept for this skill.
BANNED = re.compile(
    r'^\s*(shop\s?all|view\s?all|browse\s?all|explore\s?all|see\s?all|all\s+products?|'
    r'shop\s+the\s+collection|discover|collections?|featured|'
    r'best\s?sellers?|top\s?sellers?|sale|offers?|clearance|outlet|final\s?sale|'
    r'gift(s)?|gift\s?guides?|gift\s?cards?|'
    r'lookbooks?|blog(s)?|editorial|journal|announcements?|'
    r'magazine|inspiration|ideas|stories|landing\s?pages?|'
    r'recently\s?viewed|search\s+results?|our\s?story|about\s?us)\s*$', re.I)

# rules/new-arrivals.md SS8/SS10: "New Collection" / "New Season" / "New Range"
# is ambiguous on its own - could be a genuine new-arrivals listing or an
# editorial/campaign landing page. Force review UNLESS an unambiguous New
# Arrivals signal word is also present in the name.
NEW_ARRIVALS_STRICT = re.compile(
    r'\b(new\s?arrivals?|new\s?in|just\s?in|just\s?arrived|latest\s?arrivals?|'
    r'latest\s?products?|what.?s\s?new|new\s?this\s?season|recently\s?added|'
    r'newly\s?added)\b', re.I)
AMBIGUOUS_COLLECTION = re.compile(
    r'\b(new\s?collections?|new\s?seasons?|new\s?ranges?|new\s?edits?)\b', re.I)

MANUAL_REVIEW_COLLECTION = (
    "MANUAL REVIEW: 'COLLECTION'-STYLE NAME WITHOUT A CLEAR NEW-ARRIVALS SIGNAL "
    "- verify it is a genuine listing, not an editorial/campaign page "
    "(rules/new-arrivals.md SS8/SS10)"
)


def needs_collection_review(sub):
    if NEW_ARRIVALS_STRICT.search(sub):
        return False
    return bool(AMBIGUOUS_COLLECTION.search(sub))


# A link that resolves to a single product page, not a listing.
PDP_LIKE = re.compile(r'/(products?|item|p)/[^/]+/?$|[?&](variant|sku)=', re.I)
MANUAL_REVIEW_PDP = "MANUAL REVIEW: LINK LOOKS LIKE A SINGLE PRODUCT PAGE (PDP), NOT A PLP - verify"

# Same duplicate roster entry already resolved identically in
# Furniture/Pet Care/Seasonal/Bathroom - SR 286 = SR 165 "Noon UAE" (same
# company). Dropped entirely, not just cleared. Ledger total is intentionally
# 286 - len(this dict).
KNOWN_ROSTER_DUPLICATES = {286: 165}


def load_scope():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Output"]
    scope = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            scope.append({"sr": row[0], "company": row[1], "brand_site": row[2],
                          "country": row[3], "site_url": row[4]})
    wb.close()
    return scope


def load_blocks(scope):
    blocks, problems = [], []
    for c in scope:
        sr = c["sr"]
        p = os.path.join(SCRATCH, f"na{sr}.json")
        if not os.path.exists(p):
            blocks.append({**c, "status": "MISSING", "failure_reason":
                           f"no na{sr}.json produced", "notes": None, "rows": []})
            problems.append(f"SR {sr}: NO FILE (na{sr}.json missing)")
            continue
        try:
            d = json.load(io.open(p, encoding="utf-8"))
        except Exception as e:
            blocks.append({**c, "status": "error", "failure_reason":
                           f"worker JSON did not parse: {e}", "notes": None, "rows": []})
            problems.append(f"SR {sr}: BAD JSON: {e}")
            continue

        rows, seen_links, seen_names = [], set(), set()
        for r in d.get("rows") or []:
            sub = (r.get("sub_category") or "").strip()
            cat = (r.get("category") or "").strip() or "New Arrivals"
            link = (r.get("link") or "").strip()
            qty = r.get("qty")
            is_parent = bool(r.get("is_group") or r.get("parent"))
            flag = r.get("flag")

            if not sub:
                if cat:
                    sub = cat
                else:
                    problems.append(f"SR {sr}: row with empty sub_category dropped")
                    continue
            if BANNED.match(sub):
                problems.append(f"SR {sr}: dropped nav/marketing node '{sub}'")
                continue

            if not is_parent:
                if needs_collection_review(sub) and "COLLECTION" not in (flag or "").upper():
                    problems.append(f"SR {sr}: '{sub}' is a bare Collection-style name, no clear New Arrivals signal -> flagged")
                    flag = ((flag + " | ") if flag else "") + MANUAL_REVIEW_COLLECTION
                if link and PDP_LIKE.search(link) and "PDP" not in (flag or ""):
                    problems.append(f"SR {sr}: '{sub}' link looks like a single-product page -> flagged")
                    flag = ((flag + " | ") if flag else "") + MANUAL_REVIEW_PDP

            if is_parent:
                if qty is not None or link:
                    problems.append(f"SR {sr}: grouping row '{sub}' had qty/link -> cleared")
                qty, link = None, ""
            else:
                key = (cat.lower(), sub.lower(), link.lower())
                if key in seen_names:
                    problems.append(f"SR {sr}: dropped exact duplicate '{cat}/{sub}' -> {link}")
                    continue
                if link and link in seen_links:
                    problems.append(f"SR {sr}: '{sub}' shares a link with another row (kept, flagged)")
                    flag = ((flag + " | ") if flag else "") + \
                        "SAME LINK AS ANOTHER ROW IN THIS COMPANY - verify not an accidental duplicate"
                if qty == 0:
                    problems.append(f"SR {sr}: dropped empty (0-product) listing '{sub}'")
                    continue
                if qty is not None:
                    if not isinstance(qty, int) or isinstance(qty, bool) or qty < 0:
                        problems.append(f"SR {sr}: non-integer qty {qty!r} for '{sub}' -> nulled")
                        qty = None
                    elif not (r.get("evidence") or "").strip():
                        problems.append(f"SR {sr}: qty {qty} for '{sub}' has NO evidence -> nulled")
                        qty = None
                if link:
                    if not re.match(r'^https?://', link):
                        problems.append(f"SR {sr}: non-absolute link for '{sub}' -> nulled: {link}")
                        link = ""
                    elif re.search(r'(r\.jina\.ai|translate\.goog|web\.archive\.org|webcache)', link):
                        problems.append(f"SR {sr}: PROXY URL in link for '{sub}' -> nulled: {link}")
                        link = ""
                    elif re.search(r'[?&](q|query|search|keyword)=', link) or '/search' in link.lower():
                        problems.append(f"SR {sr}: SEARCH URL in link for '{sub}' (kept, flagged): {link}")
                        flag = ((flag + " | ") if flag else "") + \
                            "LINK LOOKS LIKE A SEARCH PAGE - verify"
                if link:
                    seen_links.add(link)
                seen_names.add(key)

            for fld, val in (("category", cat), ("sub_category", sub)):
                bad = sorted({ch for ch in val if ord(ch) > 127})
                if bad:
                    problems.append(f"SR {sr}: NON-ENGLISH chars {bad!r} in {fld} '{val}'")

            rows.append({"category": cat, "sub_category": sub, "qty": qty,
                         "link": link, "parent": is_parent,
                         "flag": flag, "evidence": r.get("evidence")})

        pruned = []
        for i, row in enumerate(rows):
            if row["parent"]:
                j = i + 1
                while j < len(rows) and rows[j]["parent"]:
                    j += 1
                if j >= len(rows):
                    problems.append(f"SR {sr}: grouping row '{row['sub_category']}' "
                                    f"has no surviving children -> dropped")
                    continue
            pruned.append(row)

        blocks.append({**c,
                       "status": d.get("status") or "unknown",
                       "failure_reason": d.get("failure_reason"),
                       "notes": d.get("notes"),
                       "rows": pruned})
    return blocks, problems


def style(ws, r, ncols=10, parent=False):
    for c in range(1, ncols + 1):
        cell = ws.cell(r, c)
        cell.font = PARENT_FONT if parent else BODY_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = CELL_BORDER
        if parent:
            cell.fill = PARENT_FILL


def ledger_status(b):
    st = (b["status"] or "").lower()
    if st in ("blocked", "failed"):
        return "BLOCKED / ACCESS FAILURE"
    if st in ("error", "missing", "unknown"):
        return "PROCESSING ERROR"
    if st == "partial" and not b["rows"]:
        # partial with zero rows means investigation was cut short before
        # anything could be confirmed either way -- not the same as a
        # completed check that found nothing. Don't call it "NO ... FOUND".
        return "BLOCKED / ACCESS FAILURE"
    return "NEW ARRIVALS FOUND" if b["rows"] else "NO NEW ARRIVALS FOUND"


def main():
    for f in ("New_Arrivals.xlsx",):
        lock = os.path.join(PROJ, "~$" + f)
        if os.path.exists(lock):
            sys.exit(f"ABORT: {f} is open in Excel (lock {lock}). Close it first.")

    scope = load_scope()
    assert len(scope) == 286, f"scope is {len(scope)}, expected 286"
    blocks, problems = load_blocks(scope)
    expected_ledger_rows = 286 - len(KNOWN_ROSTER_DUPLICATES)
    blocks = [b for b in blocks if b["sr"] not in KNOWN_ROSTER_DUPLICATES]
    for dup_sr, canonical_sr in KNOWN_ROSTER_DUPLICATES.items():
        problems.append(f"SR {dup_sr}: dropped entirely -> duplicate roster "
                         f"entry of SR {canonical_sr} (same company), not merged")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Output"

    ws.append(HDR)
    for c in range(1, 11):
        cell = ws.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 20

    r = 2
    leaf_total = parent_total = 0
    review = []
    for b in blocks:
        if not b["rows"]:
            continue
        rows = b["rows"]
        prev_cat = None
        for i, row in enumerate(rows):
            if i == 0:
                ws.cell(r, 1, b["sr"]); ws.cell(r, 2, b["company"])
                ws.cell(r, 3, b["brand_site"]); ws.cell(r, 4, b["country"])
                ws.cell(r, 5, b["site_url"]); ws.cell(r, 8, "category")
            if row["category"] and row["category"] != prev_cat:
                ws.cell(r, 6, row["category"])
                prev_cat = row["category"]
            ws.cell(r, 7, row["sub_category"])
            if row["parent"]:
                parent_total += 1
            else:
                leaf_total += 1
                if row["qty"] is not None:
                    ws.cell(r, 9, row["qty"])
                if row["link"]:
                    ws.cell(r, 10, row["link"])
                if row["qty"] is None or row.get("flag"):
                    review.append((b["sr"], b["company"], row["category"],
                                   row["sub_category"], row["qty"],
                                   row.get("flag") or "qty could not be verified",
                                   row["link"]))
            style(ws, r, parent=row["parent"])
            r += 1
        ws.cell(r, 1, None); style(ws, r); r += 1        # blank separator row

    for col, w in zip("ABCDEFGHIJ", (8, 24, 20, 10, 34, 20, 34, 10, 8, 78)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    # ---------------- Processing Ledger (all 285) ----------------
    lw = wb.create_sheet("Processing Ledger")
    LHDR = ["Maisons", "Company", "Country", "Website", "New Arrivals Found",
            "New Arrivals Record Count", "Status", "Notes"]
    lw.append(LHDR)
    for c in range(1, len(LHDR) + 1):
        cell = lw.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    counts = {"NEW ARRIVALS FOUND": 0, "NO NEW ARRIVALS FOUND": 0,
              "BLOCKED / ACCESS FAILURE": 0, "PROCESSING ERROR": 0}
    lr = 2
    for b in blocks:
        st = ledger_status(b)
        counts[st] += 1
        leaves = [x for x in b["rows"] if not x["parent"]]
        note = b.get("failure_reason") or b.get("notes") or ""
        lw.cell(lr, 1, b["sr"]); lw.cell(lr, 2, b["company"])
        lw.cell(lr, 3, b["country"]); lw.cell(lr, 4, b["site_url"])
        lw.cell(lr, 5, "YES" if b["rows"] else "NO")
        lw.cell(lr, 6, len(leaves))
        lw.cell(lr, 7, st)
        lw.cell(lr, 8, str(note)[:2000])
        style(lw, lr, ncols=len(LHDR))
        lw.cell(lr, 8).alignment = Alignment(vertical="top", wrap_text=True)
        lr += 1
    for col, w in zip("ABCDEFGH", (8, 28, 12, 40, 15, 10, 26, 100)):
        lw.column_dimensions[col].width = w
    lw.freeze_panes = "A2"

    # ---------------- Review sheet ----------------
    rw = wb.create_sheet("Review")
    RHDR = ["Maisons", "Company", "Category", "Sub-Category", "qty", "Issue", "link"]
    rw.append(RHDR)
    for c in range(1, len(RHDR) + 1):
        cell = rw.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    for i, rec in enumerate(review, start=2):
        for j, v in enumerate(rec, start=1):
            rw.cell(i, j, v)
        style(rw, i, ncols=len(RHDR))
    for col, w in zip("ABCDEFG", (8, 26, 20, 32, 8, 60, 70)):
        rw.column_dimensions[col].width = w
    rw.freeze_panes = "A2"

    wb.save(BOOK)

    # ---------------- report ----------------
    print(f"WROTE -> {BOOK}")
    print(f"  Output sheet: {r-1} rows | leaf rows {leaf_total} | grouping rows {parent_total}")
    print(f"  companies with rows: {sum(1 for b in blocks if b['rows'])}")
    print(f"  ledger rows: {lr-2}  (must be {expected_ledger_rows}: "
          f"286 roster entries minus {len(KNOWN_ROSTER_DUPLICATES)} known "
          f"roster duplicate(s))")
    print(f"  review rows: {len(review)}")
    print()
    for k, v in counts.items():
        print(f"  {k:<26} {v}")
    print(f"  {'TOTAL':<26} {sum(counts.values())}")
    if problems:
        print(f"\nDATA-HYGIENE MESSAGES ({len(problems)}):")
        for p in problems:
            print("  " + p)


if __name__ == "__main__":
    main()
