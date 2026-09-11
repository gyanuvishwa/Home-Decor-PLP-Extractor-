#!/usr/bin/env python3
"""Post-merge validation of Pet_Care.xlsx + protected-file check."""
import os, re, sys, json, hashlib, collections
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

SCRATCH = os.path.dirname(os.path.abspath(__file__))
PROJ = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
BOOK = os.path.join(PROJ, "Pet_Care.xlsx")
PROTECTED = ["Kitchen_and_Dining.xlsx", "Lighting.xlsx", "Wall_Decor.xlsx",
             "Decorative_Home_Accessories.xlsx", "Furniture.xlsx", "Textile.xlsx",
             "Storage.xlsx", "Final_Company_List (1).xlsx"]

# SR 286 is a duplicate roster entry for the identical company already
# extracted at SR 165 ("Noon UAE") -- dropped entirely by merge_pet_care.py
# per the user's explicit instruction 2026-09-02. Ledger is intentionally
# 286 - len(this set), not 286.
KNOWN_ROSTER_DUPLICATES = {286}

BANNED_URL = re.compile(r'(r\.jina\.ai|translate\.goog|web\.archive\.org|/search|[?&]q=)', re.I)
CONSUMABLE_FLAG = re.compile(
    r'\b(pet\s?)?(food|treats?|snacks?|chews?|supplements?|vitamins?|nutrition)\b', re.I)
FEEDING_EQUIPMENT_OVERRIDE = re.compile(
    r'\b(bowls?|storage|canisters?|mats?|jars?|containers?|dish(es)?|feeders?|scoops?)\b', re.I)
TOY_QUALIFIER_PHRASE = re.compile(
    r'\b(chew|treat|puzzle|interactive|fetch)?\s*toys?\b', re.I)


def is_consumable(sub):
    if FEEDING_EQUIPMENT_OVERRIDE.search(sub):
        return False
    stripped = TOY_QUALIFIER_PHRASE.sub(" ", sub)
    return bool(CONSUMABLE_FLAG.search(stripped))
MEDICAL_FLAG = re.compile(
    r'\b(medic(ine|al|ation)s?|veterinary|vet\b|prescription|clinical|'
    r'health(care)?\s?treatments?|dewormers?|flea\s?(&|and)?\s?tick)\b', re.I)


def main():
    print("=== PROTECTED FILES ===")
    base = {}
    bl = os.path.join(SCRATCH, "protected_baseline.sha256")
    if os.path.exists(bl):
        for line in open(bl, encoding="utf-8"):
            h, _, name = line.strip().partition(" ")
            base[name.lstrip("*").strip()] = h
    for f in PROTECTED:
        p = os.path.join(PROJ, f)
        if not os.path.exists(p):
            print(f"  {f:<38} (not found)")
            continue
        h = hashlib.sha256(open(p, "rb").read()).hexdigest()
        exp = base.get(f)
        state = "UNCHANGED" if exp == h else ("MODIFIED !!!" if exp else "(no baseline)")
        print(f"  {f:<38} {state}")

    print("\n=== Pet_Care.xlsx ===")
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    print("  sheets:", wb.sheetnames)
    ws = wb["Output"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    companies = [r for r in rows if r[0] is not None]
    leaves = [r for r in rows if (r[8] is not None or r[9])]
    groups = [r for r in rows if r[6] and r[8] is None and not r[9]]
    blanks = [r for r in rows if not any(v is not None for v in r)]
    print(f"  total rows          {len(rows)}")
    print(f"  companies           {len(companies)}")
    print(f"  leaf rows           {len(leaves)}")
    print(f"  grouping rows       {len(groups)}")
    print(f"  separator rows      {len(blanks)}")

    srs = [r[0] for r in companies]
    assert len(srs) == len(set(srs)), "DUPLICATE SR in Output"
    print(f"  SR unique           yes ({min(srs)}..{max(srs)})" if srs else "  (empty)")

    # duplicate (company, sub-category, qty, link)
    cur = None
    seen = collections.Counter()
    for r in rows:
        if r[0] is not None:
            cur = r[1]
        if r[9]:
            seen[(cur, r[6], r[8], r[9])] += 1
    dups = {k: v for k, v in seen.items() if v > 1}
    print(f"  duplicate records   {len(dups)}")
    for k, v in list(dups.items())[:20]:
        print("      ", v, k)

    # duplicate links within a company
    cur = None
    bycomp = collections.defaultdict(collections.Counter)
    for r in rows:
        if r[0] is not None:
            cur = r[1]
        if r[9]:
            bycomp[cur][r[9]] += 1
    dl = [(c, l, n) for c, cnt in bycomp.items() for l, n in cnt.items() if n > 1]
    print(f"  same link twice in one company: {len(dl)}")
    for x in dl[:20]:
        print("      ", x)

    # STRICT whole-file duplicate URL check -- a link must never repeat across
    # DIFFERENT companies (each URL belongs to exactly one company's domain;
    # a cross-company repeat means a copy-paste/wrong-company-assigned bug).
    cur = None
    link_owners = collections.defaultdict(set)
    link_rows = collections.defaultdict(list)
    for r in rows:
        if r[0] is not None:
            cur = (r[0], r[1])
        if r[9]:
            link_owners[r[9]].add(cur)
            link_rows[r[9]].append((cur, r[6]))
    cross_company_dupes = {link: owners for link, owners in link_owners.items() if len(owners) > 1}
    print(f"  STRICT: same URL used by >1 company: {len(cross_company_dupes)}")
    for link, owners in list(cross_company_dupes.items())[:20]:
        print("      ", link)
        for (sr, comp), sub in link_rows[link]:
            if (sr, comp) in owners:
                print("          ->", sr, comp, "/", sub)

    # non-ASCII
    na = [(r[5], r[6]) for r in rows if any(ord(ch) > 127 for ch in str(r[5] or "") + str(r[6] or ""))]
    print(f"  non-ASCII cells     {len(na)}")
    for x in na[:20]:
        print("      ", x)

    # bad URLs
    bad = [(r[6], r[9]) for r in rows if r[9] and (BANNED_URL.search(str(r[9])) or not str(r[9]).startswith("http"))]
    print(f"  suspect URLs        {len(bad)}")
    for x in bad[:20]:
        print("      ", x)

    # unresolved consumable / medical categories (rules §11/§12 - no policy yet).
    # These are expected to show up on the Review sheet with a MANUAL REVIEW flag;
    # this count is a cross-check that merge_pet_care.py's forced flag actually fired.
    cons = [(r[5], r[6]) for r in rows if r[6] and is_consumable(str(r[6]))]
    med = [(r[5], r[6]) for r in rows if r[6] and MEDICAL_FLAG.search(str(r[6]))]
    print(f"  consumable-looking sub-categories: {len(cons)}  (must all appear on Review)")
    for x in cons[:20]:
        print("      ", x)
    print(f"  medical/vet-looking sub-categories: {len(med)}  (must all appear on Review)")
    for x in med[:20]:
        print("      ", x)

    # qty sanity
    q = [r[8] for r in leaves if isinstance(r[8], int)]
    print(f"  leaf rows with qty  {len(q)} / {len(leaves)}   (null qty {len(leaves)-len(q)})")

    expected_ledger_rows = 286 - len(KNOWN_ROSTER_DUPLICATES)
    lw = wb["Processing Ledger"]
    lrows = list(lw.iter_rows(min_row=2, values_only=True))
    lrows = [r for r in lrows if r[0] is not None]
    st = collections.Counter(r[6] for r in lrows)
    print("\n=== LEDGER ===")
    print(f"  ledger rows         {len(lrows)}  (must be {expected_ledger_rows}: "
          f"286 roster entries minus {len(KNOWN_ROSTER_DUPLICATES)} known "
          f"roster duplicate(s) {sorted(KNOWN_ROSTER_DUPLICATES)})")
    lsrs = [r[0] for r in lrows]
    print(f"  unique SR           {len(set(lsrs))}")
    dup_srs_present = KNOWN_ROSTER_DUPLICATES & set(lsrs)
    print(f"  known duplicates still in ledger: {sorted(dup_srs_present) or 'none'}")
    for k, v in st.items():
        print(f"  {k:<26} {v}")
    print(f"  {'SUM':<26} {sum(st.values())}")
    assert len(lrows) == expected_ledger_rows, (
        f"LEDGER IS NOT {expected_ledger_rows} ROWS (got {len(lrows)})")
    assert not dup_srs_present, (
        f"KNOWN ROSTER DUPLICATE(S) STILL IN LEDGER: {sorted(dup_srs_present)}")

    rw = wb["Review"]
    rrows = [r for r in rw.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    print(f"\n  review rows         {len(rrows)}")

    ok = sum(st.values()) == expected_ledger_rows and not cross_company_dupes
    print("\nOK" if ok else "\nRECONCILIATION FAILED")
    assert not cross_company_dupes, (
        f"STRICT DUPLICATE URL CHECK FAILED -- {len(cross_company_dupes)} link(s) "
        "assigned to more than one company (see list above)"
    )


if __name__ == "__main__":
    main()
