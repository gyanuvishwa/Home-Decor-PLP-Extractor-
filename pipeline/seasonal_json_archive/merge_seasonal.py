#!/usr/bin/env python3
"""Build Seasonal.xlsx from the seasonal worker JSON (se<SR>.json).

Writes a BRAND NEW workbook. Never opens, reads or touches:
  Kitchen_and_Dining.xlsx / Lighting.xlsx / Wall_Decor.xlsx /
  Decorative_Home_Accessories.xlsx / Furniture.xlsx / Textile.xlsx /
  Storage.xlsx / Pet_Care.xlsx / Final_Company_List (1).xlsx  (read-only for scope)

Sheets written:
  Output            - same 10-column layout/format as the other category files
  Processing Ledger - one row per company in scope (must be 285, see
                       KNOWN_ROSTER_DUPLICATES)
  Review            - leaf rows with null qty, a MANUAL REVIEW flag, or a
                       suspect/junk-looking link

Trusts the worker's Seasonal classification (rules/seasonal.md) -- unlike
merge_furniture.py's exclusion-based NOT_FURNITURE net, Seasonal is an
inclusion-based taxonomy with no cross-category exclusivity requirement
(rules/seasonal.md SS7 -- Seasonal is explicitly allowed to overlap every
other category, so this merge does not try to net out anything another
category also claims). It nets only: navigation/marketing nodes (rules
SS4/SS17), junk/non-canonical link shapes (SS15), and the same structural
hygiene (dedup key, qty+evidence, ASCII-only) every category shares.

category = the holiday/festival name (Christmas, Diwali, Halloween, ...).
sub_category = the specific product type under it, faithful to the site's
own taxonomy (rules SS19). Do not invent a mapping between the two beyond
what the worker reported.

Idempotent: rebuilds everything from whatever se*.json is on disk.

Usage:  python merge_seasonal.py
"""
import io, json, os, re, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

SCRATCH = os.path.dirname(os.path.abspath(__file__))
PROJ = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
BOOK = os.path.join(PROJ, "Seasonal.xlsx")
SRC = os.path.join(PROJ, "Final_Company_List (1).xlsx")

HDR = ["Maisons", "Company name", "Brand Site", "Country", "site URL",
       "Category", "Sub-Category", "Type", "qty", "link"]

HDR_FONT = Font(name="Calibri", sz=10, bold=True)
HDR_FILL = PatternFill("solid", fgColor="FFD9E1F2")
PARENT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
PARENT_FONT = Font(name="Calibri", sz=10, bold=True)
BODY_FONT = Font(name="Calibri", sz=10)
THIN = Side(style="thin", color="FFBFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Navigation / marketing nodes that are never a product category
# (rules/seasonal.md SS4 and SS17, shared with the core's nav/marketing exclusions).
BANNED = re.compile(
    r'^\s*(shop\s?all|view\s?all|browse\s?all|explore\s?all|see\s?all|all\s+products?|'
    r'shop\s+the\s+collection|discover|collections?|featured|'
    r'new\s?arrivals?|new\s?in|best\s?sellers?|sale|offers?|clearance|outlet|coupons?|'
    r'gift\s?guides?|gift\s?cards?|lookbooks?|blog(s)?|editorial|journal|'
    r'magazine|inspiration|ideas|stories|landing\s?pages?|'
    r'trending|recently\s?viewed|search\s+results?|brand(s)?|designer(s)?)\s*$', re.I)

# Junk/non-canonical link shapes (rules SS15): PDP, search, tracking/campaign,
# blog/editorial, image, pagination-only URLs must never be the category link.
JUNK_URL = re.compile(
    r'(?:r\.jina\.ai|translate\.goog|web\.archive\.org|webcache|'
    r'/blog/|/blogs/|/editorial/|/journal/|/magazine/|/inspiration/|/lookbook/|'
    r'/product/|/products/[^/?]+-p-|/p/\d+|/dp/\d+|/pd/|'
    r'\.(?:jpg|jpeg|png|gif|webp|svg)(?:\?|$)|'
    r'/cart|/checkout|/account|/wishlist|'
    r'utm_[a-z]+=|[?&]campaign=|[?&]affiliate=|'
    r'[?&]page=\d+$)', re.I)
SEARCH_URL = re.compile(r'[?&](q|query|search|keyword)=|/search(/|\?|$)', re.I)

MANUAL_REVIEW_JUNK_LINK = "MANUAL REVIEW: LINK LOOKS LIKE A NON-CATEGORY (PDP/BLOG/CAMPAIGN/IMAGE) URL"

# Narrow, evidence-verified exceptions to JUNK_URL: exact URLs individually
# confirmed (orchestrator verify pass, 2026-09-04) via a live re-fetch showing
# genuine schema.org ItemList product-grid markup (numberOfItems present),
# not editorial/lookbook content. AmbienteDirect (SR 157) files its real
# holiday PLPs under /inspiration/design-special/<slug> - a site-specific URL
# quirk that the general JUNK_URL /inspiration/ heuristic correctly rejects
# for every other site but false-positives on here. Do NOT loosen the general
# pattern - extend this set instead, only after direct verification.
JUNK_URL_VERIFIED_EXCEPTIONS = {
    "https://www.ambientedirect.com/inspiration/design-special/weihnachtsdeko",
    "https://www.ambientedirect.com/inspiration/design-special/advent",
    "https://www.ambientedirect.com/inspiration/design-special/christbaumschmuck",
    "https://www.ambientedirect.com/inspiration/design-special/alessi-xmas",
    "https://www.ambientedirect.com/inspiration/design-special/valentinstag",
    "https://www.ambientedirect.com/inspiration/design-special/happy-ostern",
    "https://www.ambientedirect.com/inspiration/design-special/silvester-glamour",
    "https://www.ambientedirect.com/inspiration/design-special/geschenke-zum-vatertag",
    "https://www.ambientedirect.com/inspiration/design-special/geschenke-zum-muttertag",
}

# Same master-roster duplicate already resolved identically in Furniture and
# Pet Care (furniture_json_archive/qa_notes.md, pet_care_json_archive
# SKILL.md SS -- user decision: SR 286 = SR 165 "Noon UAE", same company).
# Dropped entirely (not even a ledger row). Ledger total is intentionally
# 286 - len(this dict).
KNOWN_ROSTER_DUPLICATES = {286: 165}


def load_scope():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Output"]
    scope = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            scope.append({"sr": row[0], "company": row[1], "brand_site": row[2],
                          "country": row[3], "site_url": row[4]})
    wb.close()
    return scope


def load_blocks(scope):
    blocks, problems = [], []
    for c in scope:
        sr = c["sr"]
        p = os.path.join(SCRATCH, f"se{sr}.json")
        if not os.path.exists(p):
            blocks.append({**c, "status": "MISSING", "failure_reason":
                           f"no se{sr}.json produced", "notes": None, "rows": []})
            problems.append(f"SR {sr}: NO FILE (se{sr}.json missing)")
            continue
        try:
            d = json.load(io.open(p, encoding="utf-8"))
        except Exception as e:
            blocks.append({**c, "status": "error", "failure_reason":
                           f"worker JSON did not parse: {e}", "notes": None, "rows": []})
            problems.append(f"SR {sr}: BAD JSON: {e}")
            continue

        rows, seen_links, seen_names = [], set(), set()
        for r in d.get("rows") or []:
            sub = (r.get("sub_category") or "").strip()
            cat = (r.get("category") or "").strip()
            link = (r.get("link") or "").strip()
            qty = r.get("qty")
            is_parent = bool(r.get("is_group") or r.get("parent"))
            flag = r.get("flag")

            if not sub:
                problems.append(f"SR {sr}: row with empty sub_category dropped")
                continue
            if not cat:
                problems.append(f"SR {sr}: row '{sub}' has empty category (holiday name) -> dropped")
                continue
            if BANNED.match(sub) or BANNED.match(cat):
                problems.append(f"SR {sr}: dropped nav/marketing node '{cat}/{sub}'")
                continue

            if is_parent:
                if qty is not None or link:
                    problems.append(f"SR {sr}: grouping row '{sub}' had qty/link -> cleared")
                qty, link = None, ""
            else:
                # A record is a duplicate only when company + category/sub-category +
                # link all repeat. A repeated NAME at a different URL is a genuinely
                # different listing page - keep both (shared merge rule).
                key = (cat.lower(), sub.lower(), link.lower())
                if key in seen_names:
                    problems.append(f"SR {sr}: dropped exact duplicate '{cat}/{sub}' -> {link}")
                    continue
                if link and link in seen_links:
                    problems.append(f"SR {sr}: '{sub}' shares a link with another row (kept, flagged)")
                    flag = ((flag + " | ") if flag else "") + \
                        "SAME LINK AS ANOTHER ROW IN THIS COMPANY - verify not an accidental duplicate"
                if qty == 0:
                    problems.append(f"SR {sr}: dropped empty (0-product) category '{sub}'")
                    continue
                if qty is not None:
                    if not isinstance(qty, int) or isinstance(qty, bool) or qty < 0:
                        problems.append(f"SR {sr}: non-integer qty {qty!r} for '{sub}' -> nulled")
                        qty = None
                    elif not (r.get("evidence") or "").strip():
                        problems.append(f"SR {sr}: qty {qty} for '{sub}' has NO evidence -> nulled")
                        qty = None
                if link:
                    if not re.match(r'^https?://', link):
                        problems.append(f"SR {sr}: non-absolute link for '{sub}' -> nulled: {link}")
                        link = ""
                    elif JUNK_URL.search(link) and link not in JUNK_URL_VERIFIED_EXCEPTIONS:
                        problems.append(f"SR {sr}: JUNK/non-category URL for '{sub}' -> nulled, flagged: {link}")
                        flag = ((flag + " | ") if flag else "") + MANUAL_REVIEW_JUNK_LINK + f" (was: {link})"
                        link = ""
                    elif SEARCH_URL.search(link):
                        problems.append(f"SR {sr}: SEARCH URL in link for '{sub}' (kept, flagged): {link}")
                        flag = ((flag + " | ") if flag else "") + \
                            "LINK LOOKS LIKE A SEARCH PAGE - verify"
                if link:
                    seen_links.add(link)
                seen_names.add(key)

            for fld, val in (("category", cat), ("sub_category", sub)):
                bad = sorted({ch for ch in val if ord(ch) > 127})
                if bad:
                    problems.append(f"SR {sr}: NON-ENGLISH chars {bad!r} in {fld} '{val}'")

            rows.append({"category": cat, "sub_category": sub, "qty": qty,
                         "link": link, "parent": is_parent,
                         "flag": flag, "evidence": r.get("evidence")})

        # Drop grouping rows left with no surviving children. Hierarchy is nested, so a
        # parent may be followed by further parents before the first leaf -- scan forward
        # over consecutive parents and require a leaf before the block ends.
        pruned = []
        for i, row in enumerate(rows):
            if row["parent"]:
                j = i + 1
                while j < len(rows) and rows[j]["parent"]:
                    j += 1
                if j >= len(rows):
                    problems.append(f"SR {sr}: grouping row '{row['sub_category']}' "
                                    f"has no surviving children -> dropped")
                    continue
            pruned.append(row)

        blocks.append({**c,
                       "status": d.get("status") or "unknown",
                       "failure_reason": d.get("failure_reason"),
                       "notes": d.get("notes"),
                       "rows": pruned})
    return blocks, problems


def style(ws, r, ncols=10, parent=False):
    for c in range(1, ncols + 1):
        cell = ws.cell(r, c)
        cell.font = PARENT_FONT if parent else BODY_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = CELL_BORDER
        if parent:
            cell.fill = PARENT_FILL


def ledger_status(b):
    """Exactly one of the shared outcomes (rules/seasonal.md SS18)."""
    st = (b["status"] or "").lower()
    if st in ("blocked", "failed"):
        return "BLOCKED / ACCESS FAILURE"
    if st == "policy_opt_out":
        return "POLICY OPT-OUT (robots.txt)"
    if st in ("error", "missing", "unknown"):
        return "PROCESSING ERROR"
    # ok / partial / found / no_seasonal
    return "SEASONAL FOUND" if b["rows"] else "NO SEASONAL FOUND"


def main():
    for f in ("Seasonal.xlsx",):
        lock = os.path.join(PROJ, "~$" + f)
        if os.path.exists(lock):
            sys.exit(f"ABORT: {f} is open in Excel (lock {lock}). Close it first.")

    scope = load_scope()
    assert len(scope) == 286, f"scope is {len(scope)}, expected 286"
    blocks, problems = load_blocks(scope)
    expected_ledger_rows = 286 - len(KNOWN_ROSTER_DUPLICATES)
    blocks = [b for b in blocks if b["sr"] not in KNOWN_ROSTER_DUPLICATES]
    for dup_sr, canonical_sr in KNOWN_ROSTER_DUPLICATES.items():
        problems.append(f"SR {dup_sr}: dropped entirely -> duplicate roster "
                         f"entry of SR {canonical_sr} (same company), not merged")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Output"

    ws.append(HDR)
    for c in range(1, 11):
        cell = ws.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 20

    r = 2
    leaf_total = parent_total = 0
    review = []
    for b in blocks:
        if not b["rows"]:
            continue
        rows = b["rows"]
        prev_cat = None
        for i, row in enumerate(rows):
            if i == 0:
                ws.cell(r, 1, b["sr"]); ws.cell(r, 2, b["company"])
                ws.cell(r, 3, b["brand_site"]); ws.cell(r, 4, b["country"])
                ws.cell(r, 5, b["site_url"]); ws.cell(r, 8, "category")
            if row["category"] and row["category"] != prev_cat:
                ws.cell(r, 6, row["category"])
                prev_cat = row["category"]
            ws.cell(r, 7, row["sub_category"])
            if row["parent"]:
                parent_total += 1
                if row.get("flag"):
                    review.append((b["sr"], b["company"], row["category"],
                                   row["sub_category"], row["qty"],
                                   row["flag"], row["link"]))
            else:
                leaf_total += 1
                if row["qty"] is not None:
                    ws.cell(r, 9, row["qty"])
                if row["link"]:
                    ws.cell(r, 10, row["link"])
                if row["qty"] is None or row.get("flag"):
                    review.append((b["sr"], b["company"], row["category"],
                                   row["sub_category"], row["qty"],
                                   row.get("flag") or "qty could not be verified",
                                   row["link"]))
            style(ws, r, parent=row["parent"])
            r += 1
        ws.cell(r, 1, None); style(ws, r); r += 1        # blank separator row

    for col, w in zip("ABCDEFGHIJ", (8, 24, 20, 10, 34, 20, 34, 10, 8, 78)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    # ---------------- Processing Ledger (285) ----------------
    lw = wb.create_sheet("Processing Ledger")
    LHDR = ["Maisons", "Company", "Country", "Website", "Seasonal Found",
            "Seasonal Record Count", "Status", "Notes"]
    lw.append(LHDR)
    for c in range(1, len(LHDR) + 1):
        cell = lw.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    counts = {"SEASONAL FOUND": 0, "NO SEASONAL FOUND": 0,
              "BLOCKED / ACCESS FAILURE": 0, "POLICY OPT-OUT (robots.txt)": 0,
              "PROCESSING ERROR": 0}
    lr = 2
    for b in blocks:
        st = ledger_status(b)
        counts[st] += 1
        leaves = [x for x in b["rows"] if not x["parent"]]
        note = b.get("failure_reason") or b.get("notes") or ""
        lw.cell(lr, 1, b["sr"]); lw.cell(lr, 2, b["company"])
        lw.cell(lr, 3, b["country"]); lw.cell(lr, 4, b["site_url"])
        lw.cell(lr, 5, "YES" if b["rows"] else "NO")
        lw.cell(lr, 6, len(leaves))
        lw.cell(lr, 7, st)
        lw.cell(lr, 8, str(note)[:2000])
        style(lw, lr, ncols=len(LHDR))
        lw.cell(lr, 8).alignment = Alignment(vertical="top", wrap_text=True)
        lr += 1
    for col, w in zip("ABCDEFGH", (8, 28, 12, 40, 15, 10, 26, 100)):
        lw.column_dimensions[col].width = w
    lw.freeze_panes = "A2"

    # ---------------- Review sheet ----------------
    rw = wb.create_sheet("Review")
    RHDR = ["Maisons", "Company", "Category", "Sub-Category", "qty", "Issue", "link"]
    rw.append(RHDR)
    for c in range(1, len(RHDR) + 1):
        cell = rw.cell(1, c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    for i, rec in enumerate(review, start=2):
        for j, v in enumerate(rec, start=1):
            rw.cell(i, j, v)
        style(rw, i, ncols=len(RHDR))
    for col, w in zip("ABCDEFG", (8, 26, 20, 32, 8, 60, 70)):
        rw.column_dimensions[col].width = w
    rw.freeze_panes = "A2"

    wb.save(BOOK)

    # ---------------- report ----------------
    print(f"WROTE -> {BOOK}")
    print(f"  Output sheet: {r-1} rows | leaf rows {leaf_total} | grouping rows {parent_total}")
    print(f"  companies with rows: {sum(1 for b in blocks if b['rows'])}")
    print(f"  ledger rows: {lr-2}  (must be {expected_ledger_rows}: "
          f"286 roster entries minus {len(KNOWN_ROSTER_DUPLICATES)} known "
          f"roster duplicate(s))")
    print(f"  review rows: {len(review)}")
    print()
    for k, v in counts.items():
        print(f"  {k:<26} {v}")
    print(f"  {'TOTAL':<26} {sum(counts.values())}")
    if problems:
        print(f"\nDATA-HYGIENE MESSAGES ({len(problems)}):")
        for p in problems:
            print("  " + p)


if __name__ == "__main__":
    main()
