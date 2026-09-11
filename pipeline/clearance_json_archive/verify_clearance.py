#!/usr/bin/env python3
"""Post-merge validation of Clearance.xlsx + protected-file check."""
import os, re, sys, json, hashlib, collections
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

SCRATCH = os.path.dirname(os.path.abspath(__file__))
PROJ = r"C:\Users\GyanendraVishwakarma\Web Research Agent"
BOOK = os.path.join(PROJ, "Clearance.xlsx")
PROTECTED = ["Kitchen_and_Dining.xlsx", "Lighting.xlsx", "Wall_Decor.xlsx",
             "Decorative_Home_Accessories.xlsx", "Furniture.xlsx", "Textile.xlsx",
             "Storage.xlsx", "Pet_Care.xlsx", "Seasonal.xlsx", "Bathroom.xlsx",
             "Final_Company_List (1).xlsx", "Bestsellers.xlsx", "New_Arrivals.xlsx"]

# SR 286 is a duplicate roster entry for the identical company already
# extracted at SR 165 ("Noon UAE") -- dropped entirely by merge_clearance.py.
# Ledger is intentionally 286 - len(this set), not 286.
KNOWN_ROSTER_DUPLICATES = {286}

BANNED_URL = re.compile(r'(r\.jina\.ai|translate\.goog|web\.archive\.org|/search|[?&]q=)', re.I)
PDP_LIKE = re.compile(r'/(products?|item|p)/[^/]+/?$|[?&](variant|sku)=', re.I)
BANNED_SUBCAT = re.compile(
    r'^\s*(shop\s?all|view\s?all|browse\s?all|explore\s?all|see\s?all|all\s+products?|'
    r'shop\s+the\s+collection|discover|collections?|featured|'
    r'new\s?arrivals?|new\s?in|just\s?in|best\s?sellers?|top\s?sellers?|'
    r'gift(s)?|gift\s?guides?|gift\s?cards?|coupons?|promo\s?codes?|'
    r'lookbooks?|blog(s)?|editorial|journal|'
    r'magazine|inspiration|ideas|stories|landing\s?pages?|black\s?friday|'
    r'recently\s?viewed|search\s+results?|our\s?story|about\s?us)\s*$', re.I)
CLEARANCE_STRICT = re.compile(
    r'\b(clearance|clearances|final\s?sale|final\s?clearance|'
    r'last\s?chance|end\s?of\s?line|discontinued)\b', re.I)
GENERIC_SALE = re.compile(
    r'\b(sale|offers?|deals?|promotions?|discounts?|special\s?offers?)\b', re.I)
# POLICY CORRECTION 2026-09-11: Outlet is NOT Clearance (rules/clearance.md
# SS1/SS3.2). No Output row's sub_category should carry "outlet" wording;
# merge_clearance.py's OUTLET_BAN drops these at merge time.
OUTLET_LEAKED = re.compile(r'\boutlet(s)?\b', re.I)
# POLICY CORRECTION 2026-09-11 (user directive): Open Box is NOT Clearance
# either; merge_clearance.py's OPEN_BOX_BAN drops these at merge time.
OPEN_BOX_LEAKED = re.compile(r'\bopen[\s-]?box\b', re.I)


def needs_sale_clearance_review(sub):
    if CLEARANCE_STRICT.search(sub):
        return False
    return bool(GENERIC_SALE.search(sub))


def main():
    print("=== PROTECTED FILES ===")
    base = {}
    bl = os.path.join(SCRATCH, "protected_baseline.sha256")
    if os.path.exists(bl):
        for line in open(bl, encoding="utf-8"):
            h, _, name = line.strip().partition(" ")
            base[name.lstrip("*").strip()] = h
    for f in PROTECTED:
        p = os.path.join(PROJ, f)
        if not os.path.exists(p):
            print(f"  {f:<38} (not found)")
            continue
        h = hashlib.sha256(open(p, "rb").read()).hexdigest()
        exp = base.get(f)
        state = "UNCHANGED" if exp == h else ("MODIFIED !!!" if exp else "(no baseline)")
        print(f"  {f:<38} {state}")

    print("\n=== Clearance.xlsx ===")
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    print("  sheets:", wb.sheetnames)
    ws = wb["Output"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    companies = [r for r in rows if r[0] is not None]
    leaves = [r for r in rows if (r[8] is not None or r[9])]
    groups = [r for r in rows if r[6] and r[8] is None and not r[9]]
    blanks = [r for r in rows if not any(v is not None for v in r)]
    print(f"  total rows          {len(rows)}")
    print(f"  companies           {len(companies)}")
    print(f"  leaf rows           {len(leaves)}")
    print(f"  grouping rows       {len(groups)}")
    print(f"  separator rows      {len(blanks)}")

    srs = [r[0] for r in companies]
    assert len(srs) == len(set(srs)), "DUPLICATE SR in Output"
    print(f"  SR unique           yes ({min(srs)}..{max(srs)})" if srs else "  (empty)")

    cur = None
    seen = collections.Counter()
    for r in rows:
        if r[0] is not None:
            cur = r[1]
        if r[9]:
            seen[(cur, r[6], r[8], r[9])] += 1
    dups = {k: v for k, v in seen.items() if v > 1}
    print(f"  duplicate records   {len(dups)}")
    for k, v in list(dups.items())[:20]:
        print("      ", v, k)

    cur = None
    bycomp = collections.defaultdict(collections.Counter)
    for r in rows:
        if r[0] is not None:
            cur = r[1]
        if r[9]:
            bycomp[cur][r[9]] += 1
    dl = [(c, l, n) for c, cnt in bycomp.items() for l, n in cnt.items() if n > 1]
    print(f"  same link twice in one company: {len(dl)}")
    for x in dl[:20]:
        print("      ", x)

    cur = None
    link_owners = collections.defaultdict(set)
    link_rows = collections.defaultdict(list)
    for r in rows:
        if r[0] is not None:
            cur = (r[0], r[1])
        if r[9]:
            link_owners[r[9]].add(cur)
            link_rows[r[9]].append((cur, r[6]))
    cross_company_dupes = {link: owners for link, owners in link_owners.items() if len(owners) > 1}
    print(f"  STRICT: same URL used by >1 company: {len(cross_company_dupes)}")
    for link, owners in list(cross_company_dupes.items())[:20]:
        print("      ", link)
        for (sr, comp), sub in link_rows[link]:
            if (sr, comp) in owners:
                print("          ->", sr, comp, "/", sub)

    na = [(r[5], r[6]) for r in rows if any(ord(ch) > 127 for ch in str(r[5] or "") + str(r[6] or ""))]
    print(f"  non-ASCII cells     {len(na)}")
    for x in na[:20]:
        print("      ", x)

    bad = [(r[6], r[9]) for r in rows if r[9] and (BANNED_URL.search(str(r[9])) or not str(r[9]).startswith("http"))]
    print(f"  suspect URLs        {len(bad)}")
    for x in bad[:20]:
        print("      ", x)

    pdp = [(r[5], r[6], r[9]) for r in rows if r[9] and PDP_LIKE.search(str(r[9]))]
    print(f"  PDP-shaped links (must all be on Review): {len(pdp)}")
    for x in pdp[:20]:
        print("      ", x)

    banned_leaked = [(r[5], r[6]) for r in rows if r[6] and BANNED_SUBCAT.match(str(r[6]))]
    print(f"  nav/marketing sub-categories leaked into Output: {len(banned_leaked)}  (must be 0)")
    for x in banned_leaked[:20]:
        print("      ", x)

    outlet_leaked = [(r[5], r[6]) for r in rows if r[6] and OUTLET_LEAKED.search(str(r[6]))]
    print(f"  Outlet-branded sub-categories leaked into Output: {len(outlet_leaked)}  (must be 0)")
    for x in outlet_leaked[:20]:
        print("      ", x)

    open_box_leaked = [(r[5], r[6]) for r in rows if r[6] and OPEN_BOX_LEAKED.search(str(r[6]))]
    print(f"  Open Box sub-categories leaked into Output: {len(open_box_leaked)}  (must be 0)")
    for x in open_box_leaked[:20]:
        print("      ", x)

    # Sale-vs-Clearance cross-check: every generic-sale-looking, non-strict
    # LEAF sub-category must appear on Review (merge_clearance.py's forced
    # flag only applies to leaf rows -- a grouping row never carries its own
    # qty/link and is never added to Review by construction, so excluding
    # grouping rows here matches actual merge behavior rather than false-
    # flagging a parent whose children were already individually vetted).
    ambiguous = [(r[5], r[6]) for r in rows
                 if r[6] and needs_sale_clearance_review(str(r[6]))
                 and not (r[8] is None and not r[9])]
    print(f"  generic Sale/Offers sub-categories not confirmed Clearance: {len(ambiguous)}  (must all appear on Review)")
    for x in ambiguous[:20]:
        print("      ", x)

    q = [r[8] for r in leaves if isinstance(r[8], int)]
    print(f"  leaf rows with qty  {len(q)} / {len(leaves)}   (null qty {len(leaves)-len(q)})")

    expected_ledger_rows = 286 - len(KNOWN_ROSTER_DUPLICATES)
    lw = wb["Processing Ledger"]
    lrows = list(lw.iter_rows(min_row=2, values_only=True))
    lrows = [r for r in lrows if r[0] is not None]
    st = collections.Counter(r[6] for r in lrows)
    print("\n=== LEDGER ===")
    print(f"  ledger rows         {len(lrows)}  (must be {expected_ledger_rows}: "
          f"286 roster entries minus {len(KNOWN_ROSTER_DUPLICATES)} known "
          f"roster duplicate(s) {sorted(KNOWN_ROSTER_DUPLICATES)})")
    lsrs = [r[0] for r in lrows]
    print(f"  unique SR           {len(set(lsrs))}")
    dup_srs_present = KNOWN_ROSTER_DUPLICATES & set(lsrs)
    print(f"  known duplicates still in ledger: {sorted(dup_srs_present) or 'none'}")
    for k, v in st.items():
        print(f"  {k:<26} {v}")
    print(f"  {'SUM':<26} {sum(st.values())}")
    assert len(lrows) == expected_ledger_rows, (
        f"LEDGER IS NOT {expected_ledger_rows} ROWS (got {len(lrows)})")
    assert not dup_srs_present, (
        f"KNOWN ROSTER DUPLICATE(S) STILL IN LEDGER: {sorted(dup_srs_present)}")

    rw = wb["Review"]
    rrows = [r for r in rw.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    print(f"\n  review rows         {len(rrows)}")

    ok = (sum(st.values()) == expected_ledger_rows and not cross_company_dupes
          and not banned_leaked and not outlet_leaked and not open_box_leaked)
    print("\nOK" if ok else "\nRECONCILIATION FAILED")
    assert not cross_company_dupes, (
        f"STRICT DUPLICATE URL CHECK FAILED -- {len(cross_company_dupes)} link(s) "
        "assigned to more than one company (see list above)"
    )
    assert not banned_leaked, (
        f"NAV/MARKETING SUB-CATEGORY LEAKED INTO OUTPUT -- {len(banned_leaked)} row(s), "
        "merge_clearance.py's BANNED filter should have caught these"
    )
    assert not outlet_leaked, (
        f"OUTLET-BRANDED SUB-CATEGORY LEAKED INTO OUTPUT -- {len(outlet_leaked)} row(s), "
        "merge_clearance.py's OUTLET_BAN filter should have caught these "
        "(policy correction 2026-09-11: Outlet is not Clearance)"
    )
    assert not open_box_leaked, (
        f"OPEN BOX SUB-CATEGORY LEAKED INTO OUTPUT -- {len(open_box_leaked)} row(s), "
        "merge_clearance.py's OPEN_BOX_BAN filter should have caught these "
        "(policy correction 2026-09-11: Open Box is not Clearance)"
    )


if __name__ == "__main__":
    main()
